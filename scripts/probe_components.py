"""Development-only component probe using newly generated, synthetic company text."""

import json
from pathlib import Path

import numpy as np
from docx import Document
from fastembed import TextEmbedding

root = Path(__file__).resolve().parents[1]
work = root / ".local" / "probe"
work.mkdir(parents=True, exist_ok=True)
source = work / "synthetic-company.docx"
doc = Document()
texts = [
    "合成测试资料：示例企业提供工业设备，客户回复应使用明确的产品型号和单位。",
    "合成测试资料：设备 AX-210 工作温度上限为 80 摄氏度，仅适用于标准配置。",
    "合成测试资料：出差报销应附发票，并由部门负责人审核。",
]
for t in texts:
    doc.add_paragraph(t)
doc.save(source)
parsed = [p.text for p in Document(source).paragraphs if p.text]
assert parsed == texts
model_name = "BAAI/bge-small-zh-v1.5"
model = TextEmbedding(model_name=model_name, cache_dir=str(work / "models"), threads=2)
vectors = np.array(list(model.embed(parsed)))
query = np.array(list(model.query_embed("AX-210 可以在多高的温度下工作？")))[0]
scores = vectors @ query / (np.linalg.norm(vectors, axis=1) * np.linalg.norm(query))
assert int(np.argmax(scores)) == 1
report = {
    "synthetic": True,
    "format": "docx",
    "paragraphs": parsed,
    "model": model_name,
    "dimension": int(vectors.shape[1]),
    "best_match": int(np.argmax(scores)),
    "scores": scores.tolist(),
}
(work / "components.json").write_text(
    json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(json.dumps({k: v for k, v in report.items() if k != "paragraphs"}, ensure_ascii=False))
