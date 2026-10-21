"""Format adapters preserve source coordinates; unavailable content is disclosed."""

import json
import math
import os
import re
import sys
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from threading import local

PARSER_VERSION = "native-v10-font-and-sheet-recovery"
CAD_LIBRARY_SUFFIXES = frozenset({".dwg", ".dxf", ".dgn", ".step", ".stp", ".iges", ".igs"})


@dataclass
class Chunk:
    text: str
    locator: dict


@dataclass
class Parsed:
    chunks: list[Chunk] = field(default_factory=list)
    status: str = "ready"
    warnings: list[str] = field(default_factory=list)

    def add(self, text, **locator):
        if "\x00" in text:
            # PDF font mappings can emit NUL separators. PostgreSQL rejects them;
            # use a space so adjacent technical tokens do not get joined together.
            text = text.replace("\x00", " ")
            note = "文本中的不可见空字符已替换为空格；原文件保留。"
            if note not in self.warnings:
                self.warnings.append(note)
        if text.strip():
            self.chunks.append(Chunk(text.strip(), locator))


def _office_guard(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 20000 or sum(i.file_size for i in entries) > 256 * 1024 * 1024:
            raise ValueError("Office 文件展开大小超过解析保护上限；原文件仍保留。")


def _table_text(rows):
    return "\n".join(" | ".join(str("" if c is None else c).replace("\n", " / ") for c in row) for row in rows)


def _word_table_rows(table):
    """python-docx repeats a horizontally merged cell at each grid position."""
    for row in table.rows:
        cells = []
        previous = None
        for cell in row.cells:
            if cell._tc is not previous:
                parts = [_word_paragraph_text(p._p) for p in cell.paragraphs]
                parts.extend(_table_text(_word_table_rows(inner)) for inner in cell.tables)
                cells.append("\n".join(part for part in parts if part.strip()))
            previous = cell._tc
        yield cells


WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _word_paragraph_text(paragraph):
    """Read Word text in order, including runs wrapped by smartTag."""
    excluded = {f"{WORD_NS}txbxContent", f"{WORD_NS}sdtContent", f"{WORD_NS}del"}
    parts = []
    for node in paragraph.iter():
        if node.tag not in {f"{WORD_NS}t", f"{WORD_NS}tab", f"{WORD_NS}br", f"{WORD_NS}cr",
                            f"{WORD_NS}footnoteReference", f"{WORD_NS}endnoteReference"}:
            continue
        if any(ancestor.tag in excluded for ancestor in node.iterancestors() if ancestor is not paragraph):
            continue
        if node.tag in {f"{WORD_NS}footnoteReference", f"{WORD_NS}endnoteReference"}:
            label = '脚注' if node.tag == f"{WORD_NS}footnoteReference" else '尾注'
            parts.append(f"[{label} {node.get(f'{WORD_NS}id', '')}]")
        elif node.tag == f"{WORD_NS}t":
            parts.append(node.text or "")
        elif node.tag == f"{WORD_NS}tab":
            parts.append("\t")
        else:
            parts.append("\n")
    return "".join(parts)


def _xml_paragraph_text(paragraph):
    return "".join(node.text or "" for node in paragraph.iter(f"{WORD_NS}t"))


def _xml_blocks(region):
    """Paragraphs and tables directly inside an XML region (text box, content control)."""
    for child in region.iterchildren():
        if child.tag == f"{WORD_NS}tbl":
            rows = []
            for row in child.findall(f"{WORD_NS}tr"):
                cells = [
                    " ".join(_xml_paragraph_text(p) for p in cell.findall(f"{WORD_NS}p")).strip()
                    for cell in row.findall(f"{WORD_NS}tc")
                ]
                if any(cells):
                    rows.append(cells)
            if rows:
                yield ("table", _table_text(rows))
        elif child.tag == f"{WORD_NS}p":
            text = _xml_paragraph_text(child).strip()
            if text:
                yield ("paragraph", text)


def _docx(path, result):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    section = ""
    paragraph = table = 0
    # Stacked and artistic titles reach us as one character per paragraph (the .doc -> .docx
    # conversion splits 方案简介 into 方/案/简/介). Held here and joined, they read as the title
    # they were; left alone they become four one-character chunks that match everything.
    pending = []
    pending_at = 0

    def flush_pending():
        if pending:
            result.add("".join(pending), kind="paragraph", paragraph=pending_at, section=section)
            pending.clear()

    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph += 1
            text = _word_paragraph_text(block._p).strip()
            if block.style and block.style.name.startswith("Heading"):
                section = text
            if len(text) <= 2 and not any(character.isdigit() for character in text):
                if not pending:
                    pending_at = paragraph
                pending.append(text)
                continue
            flush_pending()
            result.add(text, kind="paragraph", paragraph=paragraph, section=section)
        elif isinstance(block, Table):
            flush_pending()
            table += 1
            result.add(
                _table_text(_word_table_rows(block)),
                kind="table",
                table=table,
                section=section,
            )
    flush_pending()
    # Text boxes (callouts, boxed tables) and content controls (w:sdt — a generated table of
    # contents, template fields) sit outside the body traversal, so python-docx never reaches them;
    # read both from the XML instead of losing the content to a warning. A field-generated TOC is
    # how a whole table of contents went missing from a .doc while its body parsed fine.
    regions = [
        ("textbox", box, index)
        for index, box in enumerate(doc._element.xpath(".//w:txbxContent"), start=1)
    ]
    regions += [
        ("control", content, index)
        for index, content in enumerate(doc._element.xpath(".//w:sdtContent"), start=1)
    ]
    seen_regions = set()
    counted = {"textbox": 0, "control": 0}
    for kind, region, index in regions:
        blocks = list(_xml_blocks(region))
        if not blocks:
            continue
        # mc:AlternateContent stores the same region twice (Choice and Fallback); index it once.
        signature = "\n".join(text for _, text in blocks)
        if signature in seen_regions:
            continue
        seen_regions.add(signature)
        counted[kind] += 1
        for block_kind, text in blocks:
            result.add(text, kind=block_kind, **{kind: index}, section=section)
    figures = content_figure_blobs(doc)
    if figures:
        result.warnings.append(f"Word 内含 {len(figures)} 张图片／图形未解释，请查看原文件。")
    if counted["textbox"]:
        result.warnings.append(
            f"Word 文本框内容已收录（{counted['textbox']} 个文本框），版面位置与原文件不同。"
        )
    elif doc._element.xpath(".//w:txbxContent"):
        result.warnings.append("Word 文本框内只有图形，未解释，请查看原文件。")
    if counted["control"]:
        result.warnings.append(
            f"Word 内容控件内容已收录（{counted['control']} 处），版面位置可能与原文件不同。"
        )
    if doc._element.xpath(".//w:ins | .//w:del"):
        result.warnings.append(
            "Word 含修订记录，已按当前正文收录插入内容并排除删除内容；历史修订请查看原文件。"
        )
    # Headers and footers stay out of the index on purpose. Measured across the corpus they carry
    # only the company name, contact block, project line and page number — identical across dozens
    # of files — so indexing them adds repeated noise to retrieval while losing nothing. Flagging
    # them as a gap was worse than the gap: it marked 38 of 44 Word versions "partially parsed",
    # which is the signal that is supposed to mean content actually went missing.
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        from lxml import etree

        for part, tag in [('footnotes', 'footnote'), ('endnotes', 'endnote')]:
            name = f'word/{part}.xml'
            if name not in names:
                continue
            root = etree.fromstring(z.read(name), parser=etree.XMLParser(resolve_entities=False, no_network=True))
            for note in root.findall(f'{WORD_NS}{tag}'):
                if note.get(f'{WORD_NS}type') in {'separator', 'continuationSeparator', 'continuationNotice'}:
                    continue
                for kind, text in _xml_blocks(note):
                    result.add(text, kind=kind, **{tag: note.get(f'{WORD_NS}id', '')},
                               source_part=name)


EMU_PER_CM = 360000
# A picture this small is a logo, an icon or a bullet glyph — decorative, not content. Measured
# over the corpus the two groups do not overlap at all: logos and icons run 0.3–1.8 cm, content
# figures 9.4–22.7 cm. Reporting a logo as "content not indexed" made 11 versions look incomplete
# for a picture that carries no information.
FIGURE_CM = 4.0


def content_figure_blobs(document):
    """Content-size inline and floating pictures, excluding small logos and icons."""
    from docx.oxml.ns import qn

    figures = []
    for shape in document.inline_shapes:
        if max(shape.width, shape.height) / EMU_PER_CM <= FIGURE_CM:
            continue
        relation = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        part = document.part.related_parts.get(relation)
        if part is not None:
            figures.append(part.blob)
    for anchor in document._element.xpath(".//wp:anchor"):
        extent = anchor.find(qn("wp:extent"))
        if extent is None or max(int(extent.get("cx", "0")), int(extent.get("cy", "0"))) / EMU_PER_CM <= FIGURE_CM:
            continue
        for blip in anchor.xpath(".//a:blip"):
            part = document.part.related_parts.get(blip.get(qn("r:embed")))
            if part is not None:
                figures.append(part.blob)
    return figures


def _notes_have_text(archive, name):
    """True only when a footnotes/endnotes part carries an actual note.

    Word and LibreOffice always write the part — a document with no footnotes still has
    `word/footnotes.xml` holding just the separator and continuationSeparator entries. Warning on
    the part's presence marked 30 of 54 versions "partially parsed" for nothing, the same way the
    header/footer warning did.
    """
    xml = archive.read(name).decode("utf-8", errors="replace")
    structural = re.compile(
        r"<w:(?:foot|end)note\b[^>]*w:type=\"(?:separator|continuationSeparator|continuationNotice)\"[^>]*>.*?</w:(?:foot|end)note>",
        re.DOTALL,
    )
    return bool(re.search(r"<w:t(?:\s[^>]*)?>\s*\S", structural.sub("", xml)))


def _xlsx_content_bounds(path):
    """Ignore dimension/style-only cells without changing the uploaded workbook."""
    import posixpath
    from xml.etree import ElementTree as ET

    from openpyxl.utils.cell import coordinate_to_tuple

    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    bounds = {}
    with zipfile.ZipFile(path) as archive:
        relationships = {
            node.get("Id"): node.get("Target")
            for node in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            if node.get("TargetMode") != "External"
        }
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        for sheet in workbook.findall(f"{ns}sheets/{ns}sheet"):
            target = relationships[sheet.get(f"{rel_ns}id")]
            member = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            max_row = max_column = 0
            with archive.open(member) as source:
                for _, element in ET.iterparse(source, events=("end",)):
                    if element.tag == f"{ns}c":
                        value = element.find(f"{ns}v")
                        inline = element.find(f"{ns}is")
                        formula = element.find(f"{ns}f")
                        populated = formula is not None or (value is not None and value.text is not None)
                        populated = populated or (inline is not None and any(t.text for t in inline.iter(f"{ns}t")))
                        if populated:
                            row, column = coordinate_to_tuple(element.attrib["r"])
                            max_row, max_column = max(max_row, row), max(max_column, column)
                        element.clear()
                    elif element.tag == f"{ns}row":
                        element.clear()
            bounds[sheet.get("name")] = (max_row, max_column)
    return bounds


def _sheet_rows(sheet, bounds=None):
    max_row, max_column = bounds if bounds is not None else (sheet.max_row or 0, sheet.max_column or 0)
    if bounds is not None and (not max_row or not max_column):
        return []
    if max_row * max_column > 1_000_000:
        raise ValueError("工作表范围超过解析保护上限；原文件仍保留。")
    rows = []
    width = 0
    # The XLSX dimension element is optional. Bound actual rows before accumulating,
    # including unsized sheets and sparse coordinates; don't rely on cached dimensions.
    for index, row in enumerate(sheet.iter_rows(max_row=max_row or None, max_col=max_column or None, values_only=True), 1):
        width = max(width, len(row))
        if index * max(1, width) > 1_000_000:
            raise ValueError("工作表范围超过解析保护上限；原文件仍保留。")
        rows.append(row)
    return rows


def _xlsx(path, result):
    from contextlib import ExitStack

    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    from .excel_visuals import extract_text_shapes, rich_picture_cells, visual_counts

    rich_cells = rich_picture_cells(path)
    content_bounds = _xlsx_content_bounds(path)

    with ExitStack() as resources:
        # Storage uses opaque .blob names; the validated original filename chooses
        # this adapter. File handles avoid openpyxl's path-extension check.
        formulas = load_workbook(
            resources.enter_context(path.open("rb")), read_only=True, data_only=False
        )
        resources.callback(formulas.close)
        cached = load_workbook(
            resources.enter_context(path.open("rb")), read_only=True, data_only=True
        )
        resources.callback(cached.close)
        for sheet in formulas:
            bounds = content_bounds[sheet.title]
            values = _sheet_rows(cached[sheet.title], bounds)
            source_rows = _sheet_rows(sheet, bounds)
            columns = sorted({c for row in source_rows for c, value in enumerate(row) if value is not None and value != ""})
            rows = []
            row_numbers = []
            has_formula = False
            missing_formula_cache = False
            for r, row in enumerate(source_rows):
                if not any(value is not None and value != "" for value in row):
                    continue
                output = []
                for c in columns:
                    value = row[c] if c < len(row) else None
                    if (sheet.title, f"{get_column_letter(c + 1)}{r + 1}") in rich_cells and value == "#VALUE!":
                        value = ""
                    if isinstance(value, str) and value.startswith("="):
                        has_formula = True
                        saved = values[r][c] if r < len(values) and c < len(values[r]) else None
                        if saved is None:
                            missing_formula_cache = True
                        value = (
                            f"{value} [缓存结果: {saved}]"
                            if saved is not None
                            else f"{value} [缓存结果缺失]"
                        )
                    output.append("" if value is None else str(value))
                rows.append(output)
                row_numbers.append(r + 1)
            # Keep each small sheet whole; large sheets repeat their header across row groups.
            for start in range(0, len(rows), 50):
                group = rows[start : start + 50]
                if start and rows:
                    group = [rows[0], *group]
                result.add(
                    _table_text(group),
                    kind="table",
                    sheet=sheet.title,
                    range=f"{get_column_letter(columns[0] + 1)}{row_numbers[start]}:{get_column_letter(columns[-1] + 1)}{row_numbers[min(start + 50, len(rows)) - 1]}",
                    columns=[get_column_letter(c + 1) for c in columns],
                    source_rows=row_numbers[start:start + 50],
                    header_row=row_numbers[0],
                    repeated_header="1" if start else None,
                )
            if has_formula:
                result.warnings.append(
                    f"工作表“{sheet.title}”公式结果未解析（原文件无缓存），请核对原文件。"
                    if missing_formula_cache else
                    f"工作表“{sheet.title}”保留公式与已有缓存结果，未重新计算公式。"
                )
        visuals = visual_counts(path)
        if visuals["pictures"]:
            result.warnings.append(
                f"Excel 内含 {visuals['pictures']} 张图片未解释，请查看原文件。"
            )
        if visuals["charts"]:
            result.warnings.append("Excel 图表未解析，请查看原文件。")
        if visuals["text_shapes"]:
            shapes = extract_text_shapes(path)
            for shape in shapes:
                result.add(shape['text'], kind='drawing_text', sheet=shape['sheet'],
                           cell=shape['cell'], source_part=shape['source_part'])
            if len(shapes) < visuals['text_shapes']:
                result.warnings.append("Excel 部分绘图文字未解析，请查看原文件。")


def _ppt_has_visual(shape, slide_area):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    area = shape.width * shape.height / slide_area if slide_area else 0
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return area >= 0.05 or any(_ppt_has_visual(child, slide_area) for child in shape.shapes)
    if shape.has_table or shape.has_text_frame:
        return False
    if shape.shape_type in {
        MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.CHART,
        MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
    }:
        return area >= 0.01
    return shape.shape_type == MSO_SHAPE_TYPE.FREEFORM and area >= 0.05


def ppt_visual_slides(path):
    from pptx import Presentation

    deck = Presentation(path)
    area = deck.slide_width * deck.slide_height
    return [index for index, slide in enumerate(deck.slides, 1)
            if any(_ppt_has_visual(shape, area) for shape in slide.shapes)]


def _pptx(path, result):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(path)
    slide_area = deck.slide_width * deck.slide_height
    for index, slide in enumerate(deck.slides, 1):
        visual = False

        def visit(shapes, slide_number, prefix=""):
            nonlocal visual
            for shape_index, shape in enumerate(shapes, 1):
                locator = f"{prefix}{shape_index}" if prefix else shape_index
                if _ppt_has_visual(shape, slide_area):
                    visual = True
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    visit(shape.shapes, slide_number, f"{locator}.")
                elif shape.has_table:
                    result.add(
                        _table_text([[c.text for c in row.cells] for row in shape.table.rows]),
                        kind="table", slide=slide_number, table=locator,
                    )
                elif shape.has_text_frame:
                    result.add(shape.text, kind="slide_text", slide=slide_number, shape=locator)

        visit(slide.shapes, index)
        if visual:
            result.warnings.append(f"幻灯片 {index} 含未解释的图形／图片，需查看原文件。")
        if slide.has_notes_slide:
            text = slide.notes_slide.notes_text_frame.text
            result.add(text, kind="slide_notes", slide=index)


_ocr_thread = local()


def _ocr_engine():
    if hasattr(_ocr_thread, "engine"):
        return _ocr_thread.engine
    from rapidocr import RapidOCR

    params = {
        "EngineConfig.onnxruntime.intra_op_num_threads": 2,
        "EngineConfig.onnxruntime.inter_op_num_threads": 2,
    }
    if cache := os.environ.get("AGICO_KB_MODEL_CACHE"):
        params["Global.model_root_dir"] = str(Path(cache) / "ocr")
    _ocr_thread.engine = RapidOCR(params=params)
    return _ocr_thread.engine


def _ocr_rect(box):
    xs = [float(point[0]) for point in box]
    ys = [float(point[1]) for point in box]
    return min(xs), min(ys), max(xs), max(ys)


def _ocr_gaps(intervals, minimum):
    """Split points where whitespace wider than `minimum` separates covered spans."""
    ordered = sorted(intervals)
    found = []
    reach = ordered[0][1]
    for start, end in ordered[1:]:
        if start - reach > minimum:
            found.append((reach, start))
        reach = max(reach, end)
    return found


def _ocr_blocks(items, width, height, depth=0):
    """Recursive XY-cut: split a region on whitespace that spans it entirely, so multi-column
    and staggered layouts (timelines, radial charts) do not interleave their text."""
    if len(items) <= 1 or depth > 8:
        return [items]
    for axis in ("v", "h"):
        if axis == "v":
            found = _ocr_gaps([(item[0], item[2]) for item in items], max(14.0, width * 0.02))
        else:
            found = _ocr_gaps([(item[1], item[3]) for item in items], max(10.0, height * 0.015))
        if not found:
            continue
        cut = (found[0][0] + found[0][1]) / 2
        index = 0 if axis == "v" else 1
        before = [i for i in items if (i[index] + i[index + 2]) / 2 < cut]
        after = [i for i in items if (i[index] + i[index + 2]) / 2 >= cut]
        if not before or not after:
            continue
        return _ocr_blocks(before, width, height, depth + 1) + _ocr_blocks(
            after, width, height, depth + 1
        )
    return [items]


def _ocr_block_text(block):
    """Order one block into lines, stitching lines the layout wrapped mid-word."""
    lines = []
    for item in sorted(block, key=lambda i: (i[1], i[0])):
        for line in lines:
            overlap = min(line["y1"], item[3]) - max(line["y0"], item[1])
            if overlap > 0.5 * min(line["y1"] - line["y0"], item[3] - item[1]):
                line["items"].append(item)
                line["y0"] = min(line["y0"], item[1])
                line["y1"] = max(line["y1"], item[3])
                break
        else:
            lines.append({"y0": item[1], "y1": item[3], "items": [item]})
    lines.sort(key=lambda line: line["y0"])
    left = min(item[0] for item in block)
    right = max(item[2] for item in block)
    margin = 0.12 * (right - left)
    text = ""
    for index, line in enumerate(lines):
        line["items"].sort(key=lambda i: i[0])
        joined = "".join(item[4] for item in line["items"])
        if index == 0:
            text = joined
            continue
        previous = lines[index - 1]
        wrapped = (
            max(item[2] for item in previous["items"]) >= right - margin
            and min(item[0] for item in line["items"]) <= left + margin
        )
        text += ("" if wrapped else "\n") + joined
    return text


def _ocr_reading_order(boxes, texts):
    """Rebuild OCR reading order from detection geometry instead of trusting detector order."""
    items = []
    for box, text in zip(boxes, texts):
        x0, y0, x1, y1 = _ocr_rect(box)
        items.append((x0, y0, x1, y1, text))
    if not items:
        return "\n".join(texts)
    width = max(item[2] for item in items)
    height = max(item[3] for item in items)
    blocks = _ocr_blocks(items, width, height)
    return "\n".join(_ocr_block_text(block) for block in blocks if block)


def _squash(text):
    """Whitespace-insensitive key, so OCR and native text of the same line compare equal."""
    return "".join(text.split())


# A folio: "3", "- 3 -", "第 3 页", "3 / 10", "Page 3 of 10". A trailing dot is excluded on
# purpose — "2." at the top of a page is a list number, not a page number.
PAGE_NUMBER = re.compile(
    r"^[\s\-–—_·/]*(?:第\s*)?(?:page\s*)?\d{1,4}(?:\s*(?:/|of|共)\s*\d{1,4})?\s*(?:页)?[\s\-–—_·/]*$",
    re.IGNORECASE,
)


def _page_lines(page):
    """Text lines with their vertical centre as a fraction of page height."""
    height = page.rect.height or 1.0
    lines = []
    import pymupdf

    for block in page.get_text("dict", flags=pymupdf.TEXTFLAGS_DICT & ~pymupdf.TEXT_PRESERVE_IMAGES)["blocks"]:
        if block.get("type") != 0:
            continue
        for line in block["lines"]:
            text = "".join(span["text"] for span in line["spans"]).strip()
            if text:
                lines.append((text, (line["bbox"][1] + line["bbox"][3]) / 2 / height))
    return lines


def _in_margin(position):
    """True for a line sitting in the top or bottom band of the page — where furniture lives."""
    return position <= 0.12 or position >= 0.88


def _furniture(pages_lines):
    """Running heads and footers: short text repeated in the top or bottom band of most pages.

    They are the largest single source of retrieval noise in a PDF — a footer's phone number gets
    indexed once per page — while carrying almost nothing: page numbers, company name, contact
    block. Dropping them is safe; keeping them is not. Position is part of the test on purpose: a
    table header row repeated in the middle of a page is content and stays, and a dense document
    whose body text runs into the margins keeps that text because it does not repeat.
    """
    pages_lines = list(pages_lines)
    if len(pages_lines) < 3:
        return frozenset()
    threshold = max(3, math.ceil(0.6 * len(pages_lines)))
    seen = {}
    folio_pages = set()
    for index, lines in enumerate(pages_lines):
        for text, position in lines:
            if not _in_margin(position):
                continue
            if PAGE_NUMBER.match(text):
                folio_pages.add(index)
            elif len(text) <= 120:
                seen.setdefault(_squash(text), set()).add(index)
    drop = {key for key, pages in seen.items() if len(pages) >= threshold}
    if len(folio_pages) >= threshold:
        # A folio differs on every page, so repetition can never see it; left alone it becomes a
        # chunk whose entire text is "3". Require the pattern on most pages before dropping, so a
        # lone number at the bottom of one page is treated as content.
        for lines in pages_lines:
            for text, position in lines:
                if _in_margin(position) and PAGE_NUMBER.match(text):
                    drop.add(_squash(text))
    return frozenset(drop)


def _drop_furniture(text, drop):
    """Same text, minus the lines already identified as running heads or footers."""
    if not drop:
        return text
    kept = [line for line in text.splitlines() if _squash(line) not in drop]
    return "\n".join(kept)


def _ocr(image_input, result, drop=frozenset(), **locator):
    output = _ocr_engine()(image_input)
    assessment = {}
    if (locator.get('page') and not locator.get('tile')) or locator.get('image'):
        from .pdf_screening import text_scan_assessment
        assessment = text_scan_assessment(
            image_input, getattr(output, 'boxes', None), getattr(output, 'txts', None),
            getattr(output, 'scores', None),
        )
    texts = getattr(output, "txts", None)
    if texts:
        boxes = getattr(output, "boxes", None)
        if boxes is not None and len(boxes) == len(texts):
            keep = [index for index, text in enumerate(texts) if _squash(text) not in drop]
            texts = [texts[index] for index in keep]
            boxes = [boxes[index] for index in keep]
        else:
            texts = [text for text in texts if _squash(text) not in drop]
            boxes = None
        try:
            text = (
                _ocr_reading_order(boxes, texts)
                if boxes is not None and len(boxes)
                else "\n".join(texts)
            )
        except (AttributeError, IndexError, TypeError, ValueError):
            # Geometry from an unexpected detector build must never lose the recognised text.
            text = "\n".join(texts)
        if text.strip():
            assessment_key = "image_assessment" if locator.get('image') else "pdf_assessment"
            if locator.get('page') and boxes is not None:
                from .drawing_index import assess

                height, width = image_input.shape[:2]
                clip = locator.get('region', [0, 0, width, height])
                records = []
                scores = getattr(output, 'scores', None)
                for ri, (box, line) in enumerate(zip(boxes, texts)):
                    x0, y0, x1, y1 = _ocr_rect(box)
                    record = {'text': line, 'region': [
                        clip[0] + x0 / width * (clip[2] - clip[0]),
                        clip[1] + y0 / height * (clip[3] - clip[1]),
                        clip[0] + x1 / width * (clip[2] - clip[0]),
                        clip[1] + y1 / height * (clip[3] - clip[1])],
                        'group': ri, 'source': 'ocr',
                        'confidence': float(scores[keep[ri]]) if scores is not None else 0.0}
                    records.append(record)
                locator['drawing_records'] = records
                assessment['drawing_preflight'] = assess(records, clip[2]-clip[0], clip[3]-clip[1], scan=True)
            result.add(text, kind="ocr", ocr=True, **locator, **{assessment_key: assessment})
    result.warnings.append(
        "OCR 识别文字未经人工核对；表格结构与图形语义可能不完整，关键数字请核对原文件。"
    )


def _ocr_page(page, result, index, drop=frozenset()):
    """Render bounded tiles so large engineering sheets retain readable labels."""
    import numpy as np
    import pymupdf

    rect = page.rect
    dpi = 200 if max(rect.width, rect.height) > 1200 else 300
    scale = min(dpi / 72, math.sqrt(64 * 1900**2 / max(1, rect.width * rect.height)))
    while math.ceil(rect.width * scale / 1900) * math.ceil(rect.height * scale / 1900) > 64:
        scale *= 0.95
    width, height = math.ceil(rect.width * scale), math.ceil(rect.height * scale)
    if max(width, height) <= 4000 or (max(rect.width, rect.height) <= 1100 and max(width, height) <= 4500):
        pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale),
                              colorspace=pymupdf.csRGB, alpha=False)
        _ocr(np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, 3),
             result, drop=drop, page=index, region=[round(v, 2) for v in rect])
        return
    tile_size, stride = 2000, 1900
    tile = 0
    for y in range(0, height, stride):
        for x in range(0, width, stride):
            clip = pymupdf.Rect(rect.x0 + x / scale, rect.y0 + y / scale,
                               rect.x0 + min(x + tile_size, width) / scale,
                               rect.y0 + min(y + tile_size, height) / scale)
            pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), clip=clip,
                                  colorspace=pymupdf.csRGB, alpha=False)
            array = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, 3)
            if np.mean(np.min(array, axis=2) < 220) < 0.0004:
                continue
            tile += 1
            _ocr(array, result, drop=drop, page=index, tile=tile,
                 region=[round(v, 2) for v in clip])


def _vector_content(page):
    """Complex vector paths identify illustrations and outlined CAD text, not table grids."""
    drawings = page.get_drawings()
    complex_paths = sum(
        len(path['items']) > 2 or any(item[0] == 'c' for item in path['items'])
        for path in drawings
    )
    return complex_paths >= 40


def _image_coverage(page):
    """Fraction of the page area covered by placed images (0..1), used to spot picture pages."""
    import pymupdf

    rect = page.rect
    area = abs(rect.width * rect.height) or 1.0
    covered = 0.0
    for item in page.get_image_info():
        box = pymupdf.Rect(item["bbox"]).intersect(rect)
        if not box.is_empty:
            covered += abs(box.width * box.height)
    return min(1.0, covered / area)


def _content_images(page):
    """Placed images big enough to be content rather than decoration.

    A logo or stamp in the margin, or anything under 1% of the page area, carries no information:
    flagging it made a page look like it hid content when it did not.
    """
    import pymupdf

    rect = page.rect
    area = abs(rect.width * rect.height) or 1.0
    found = []
    for item in page.get_image_info():
        box = pymupdf.Rect(item["bbox"])
        if box.is_empty:
            continue
        share = abs(box.width * box.height) / area
        if share < 0.01:
            continue
        centre = (box.y0 + box.y1) / 2 / (rect.height or 1.0)
        if share < 0.05 and _in_margin(centre):
            continue
        found.append(box)
    return found


def _pdf_scan_records(page, index):
    """Run local OCR once and retain page coordinates for scan classification/indexing."""
    scanned = Parsed()
    _ocr_page(page, scanned, index, drop=frozenset())
    records = []
    for ci, chunk in enumerate(scanned.chunks):
        for record in chunk.locator.pop('drawing_records', []):
            record['group'] = ('ocr', ci, record['group'])
            records.append(record)
    return scanned, records


def _pdf(path, result, first_page=None, last_page=None):
    """Parse a PDF. first_page/last_page (1-based, inclusive) bound the work so the worker can
    split huge scans into segments that each fit the parse timeout."""
    import pymupdf

    from .drawing_index import add_index, broken_mapping, native_records, preflight

    with pymupdf.open(path) as doc:
        if doc.needs_pass:
            raise ValueError("PDF 已加密，需提供可读取版本。")
        start = max(1, first_page or 1)
        end = min(doc.page_count, last_page or doc.page_count)
        pages = [doc[index - 1] for index in range(start, end + 1)]
        # Running heads and footers repeat on every page; dropping them is the difference between
        # one phone number in the index and one per page. The page text layer of a picture-dominant
        # document is nothing but this furniture, so the same keys also filter the OCR output.
        drop = _furniture(_page_lines(page) for page in pages)
        for index, page in zip(range(start, end + 1), pages):
            native = native_records(page)
            font_recovery = any(broken_mapping(r['text']) for r in native)
            # Once this page's font mapping is proven corrupt, short fragments
            # from the same anomalous glyph range are unreliable too. The OCR
            # path reads their pixels rather than dropping their visible content.
            trusted = [r for r in native if not broken_mapping(r['text']) and not (
                font_recovery and any('\u0180' <= c <= '\u024f' for c in r['text']))]
            assessment = preflight(page, trusted)
            if assessment['content_type'] == 'drawing':
                records = trusted
                if _image_coverage(page) >= .2 or font_recovery:
                    _, raster_records = _pdf_scan_records(page, index)
                    # A translated/native title does not prove the scanned notes
                    # and parts list were already read. Keep that raster evidence.
                    native_text = {_squash(r['text']) for r in records}
                    records.extend(r for r in raster_records if _squash(r['text']) not in native_text)
                add_index(result, records, index, assessment)
                continue
            text = _drop_furniture(str(page.get_text()), drop).strip()
            visible_chars = sum(not ch.isspace() and ch != "\x00" for ch in text)
            # An image-dominant page (slide decks exported as pictures, full-bleed scans) still
            # carries a thin text layer: the header/footer furniture. Judging by "is there any
            # text" would skip the whole document, so judge by how little text there is against
            # how much of the page is picture.
            coverage = _image_coverage(page)
            vector_content = _vector_content(page)
            image_dominant = bool(page.get_images()) and coverage >= 0.2
            if font_recovery or (vector_content and visible_chars < 400) or (visible_chars < 15 and page.get_images()) or image_dominant:
                # OCR is the inexpensive scan preflight and is reused by the full
                # document path. Classify the entire page, not independent tiles.
                scanned, records = _pdf_scan_records(page, index)
                scanned_assessment = preflight(page, records) if records else assessment
                if scanned_assessment['content_type'] == 'drawing':
                    add_index(result, records, index, scanned_assessment)
                    continue
                for chunk in scanned.chunks:
                    chunk.text = _drop_furniture(chunk.text, drop)
                    if font_recovery:
                        chunk.locator['font_recovery'] = True
                result.chunks.extend(c for c in scanned.chunks if c.text.strip())
                result.warnings.extend(scanned.warnings)
                if font_recovery:
                    # Pixels, rather than the corrupt font mapping, are now the
                    # source. Do not index a second contradictory garbled layer.
                    if any(c.locator.get('pdf_assessment', {}).get('route') != 'ocr_text'
                           for c in scanned.chunks) or not scanned.chunks:
                        result.warnings.append(
                            f'PDF 第 {index} 页字体映射异常，已转为图形识别；内容尚待视觉核对。')
                    continue
            if vector_content and visible_chars < 400:
                if text:
                    result.add(text, kind="page", page=index)
                result.warnings.append(
                    f"PDF 第 {index} 页含矢量图形／轮廓文字，已 OCR 识别；图形内容未解释。"
                )
                continue
            if (visible_chars < 15 and page.get_images()) or image_dominant:
                if text:
                    result.add(text, kind="page", page=index)
                if visible_chars < 15:
                    result.warnings.append(
                        f"PDF 第 {index} 页为扫描页，已 OCR 识别；关键数字请核对原文件。"
                    )
                else:
                    result.warnings.append(
                        f"PDF 第 {index} 页以图片为主（图片覆盖 {coverage * 100:.0f}%），"
                        f"已对整页做 OCR；图形语义仍需核对原文件。"
                    )
                continue
            # Dense CAD lines are not spreadsheet cell borders. Table detection can treat
            # the whole drawing as one cell grid and discard labels inside its bounds.
            drawing_sheet = vector_content and max(page.rect.width, page.rect.height) > 1200
            tables = [] if drawing_sheet else page.find_tables().tables
            for ti, table in enumerate(tables, 1):
                result.add(_table_text(table.extract()), kind="table", page=index, table=ti)
            blocks = []
            for block in page.get_text("blocks", sort=True):
                if block[6] == 0 and not any(
                    pymupdf.Rect(t.bbox).contains(pymupdf.Rect(block[:4])) for t in tables
                ):
                    body = _drop_furniture(block[4], drop).strip()
                    if body:
                        if drawing_sheet:
                            result.add(body, kind="page", page=index,
                                       region=[round(v, 2) for v in block[:4]])
                        else:
                            blocks.append(body)
            if blocks:
                result.add("\n".join(blocks), kind="page", page=index)
            if vector_content:
                result.warnings.append(f"PDF 第 {index} 页含矢量图形，原生文字已提取，图形内容未解释。")
            elif _content_images(page):
                result.warnings.append(f"PDF 第 {index} 页含图片，原生文字已提取，图片内容未解释。")


def parse_file(path: Path, filename: str, first_page=None, last_page=None) -> Parsed:
    result = Parsed()
    suffix = Path(filename).suffix.lower()
    if suffix in {".docx", ".xlsx", ".pptx"}:
        _office_guard(path)
    adapters = {".docx": _docx, ".xlsx": _xlsx, ".pptx": _pptx}
    if suffix == ".pdf":
        _pdf(path, result, first_page=first_page, last_page=last_page)
    elif suffix in adapters:
        adapters[suffix](path, result)
    elif suffix in {".xls", ".doc", ".ppt"}:
        # Pure-Python recovery path (LibreOffice absence or failure): values over fidelity.
        from .fallback_parse import FALLBACKS

        FALLBACKS[suffix](path, result)
    elif suffix in {".md", ".txt", ".csv", ".json"}:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("gb18030")
            result.warnings.append("文件按 GB18030 解码，请核对中文。")
        section = ""
        for index, block in enumerate(text.split("\n\n"), 1):
            if block.startswith("#"):
                section = block.splitlines()[0].lstrip("# ").strip()
            result.add(block, kind="text", paragraph=index, section=section)
    elif suffix in {".dwg", ".dxf"}:
        from .cad import parse_cad

        return parse_cad(path, suffix, result)
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        import numpy as np
        from PIL import Image, ImageOps

        with Image.open(path) as image:
            pixels = np.asarray(ImageOps.exif_transpose(image).convert("RGB"))
            _ocr(pixels, result, image=1)
        if not result.chunks:
            result.status = "stored_only"
            result.warnings = ["图片未检测到可辨认文字，按规则跳过转述，原文件保留。"]
            return result
    else:
        result.status = "stored_only"
        result.warnings.append(
            "此格式仅保存原文件，暂无正文索引；旧 Office 请另提交 .docx/.xlsx/.pptx，CAD／模型包请关联说明或参数文件。"
        )
        return result
    result.warnings = list(dict.fromkeys(result.warnings))
    if not result.chunks:
        result.status = "failed"
        result.warnings.append("未提取到可读取文字；原文件保留。")
    elif any(
        any(marker in w for marker in ["未解析", "未完整解析", "未解释", "OCR", "未纳入"])
        for w in result.warnings
    ):
        result.status = "partial"
    # Keep every actual gap for API/worker decisions. Presentation handles routine
    # notes separately; truncation used to discard later-page failures silently.
    return result


if __name__ == "__main__":
    # The worker invokes this in a bounded subprocess, with server-generated paths only.
    # Optional argv[4]/argv[5] bound the page range for segmented PDF parsing.
    first = int(sys.argv[4]) if len(sys.argv) > 4 and sys.argv[4] else None
    last = int(sys.argv[5]) if len(sys.argv) > 5 and sys.argv[5] else None
    parsed = parse_file(Path(sys.argv[1]), sys.argv[2], first_page=first, last_page=last)
    Path(sys.argv[3]).write_text(json.dumps(asdict(parsed), ensure_ascii=False), encoding="utf-8")
