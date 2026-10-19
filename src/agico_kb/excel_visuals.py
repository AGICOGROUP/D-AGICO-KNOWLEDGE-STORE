"""Locate worksheet and in-cell pictures without treating empty drawing shapes as content."""

import posixpath
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET


@dataclass
class ExcelPicture:
    sheet: str
    cell: str
    blob: bytes


def _part(base: str, target: str) -> str:
    return target.lstrip("/") if target.startswith("/") else posixpath.normpath(
        posixpath.join(posixpath.dirname(base), target)
    )


def _rich_pictures(archive: zipfile.ZipFile) -> list[ExcelPicture]:
    """Resolve Excel's newer in-cell images through value metadata and richData rels."""
    names = set(archive.namelist())
    required = {
        "xl/metadata.xml", "xl/richData/rdrichvalue.xml",
        "xl/richData/rdrichvaluestructure.xml", "xl/richData/richValueRel.xml",
        "xl/richData/_rels/richValueRel.xml.rels", "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
    }
    if not required.issubset(names):
        return []
    metadata = ET.fromstring(archive.read("xl/metadata.xml"))
    value = metadata.find("{*}valueMetadata")
    future = metadata.find("{*}futureMetadata")
    if value is None or future is None:
        return []
    value_bks = value.findall("{*}bk")
    future_bks = future.findall("{*}bk")
    rich_values = ET.fromstring(archive.read("xl/richData/rdrichvalue.xml")).findall("{*}rv")
    structures = ET.fromstring(
        archive.read("xl/richData/rdrichvaluestructure.xml")
    ).findall("{*}s")
    rich_rels = ET.fromstring(archive.read("xl/richData/richValueRel.xml"))
    ids = [next((value for key, value in rel.attrib.items() if key.endswith("}id")), "")
           for rel in rich_rels.findall("{*}rel")]
    targets = {
        rel.get("Id"): _part("xl/richData/richValueRel.xml", rel.get("Target", ""))
        for rel in ET.fromstring(archive.read("xl/richData/_rels/richValueRel.xml.rels"))
    }
    sheet_targets = {
        rel.get("Id"): _part("xl/workbook.xml", rel.get("Target", ""))
        for rel in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    }
    sheets = {}
    for sheet in ET.fromstring(archive.read("xl/workbook.xml")).findall(".//{*}sheet"):
        relation = next((v for k, v in sheet.attrib.items() if k.endswith("}id")), "")
        if relation in sheet_targets:
            sheets[sheet_targets[relation]] = sheet.get("name", "")
    found = []
    for part, sheet_name in sheets.items():
        if part not in names:
            continue
        root = ET.fromstring(archive.read(part))
        for cell in root.findall(".//{*}c"):
            try:
                vm = int(cell.get("vm", "0")) - 1
                if vm < 0:
                    continue
                rc = value_bks[vm].find("{*}rc")
                future_index = int(rc.get("v"))
                rvb = future_bks[future_index].find(".//{*}rvb")
                rv = rich_values[int(rvb.get("i"))]
                structure = structures[int(rv.get("s", "0"))]
                if structure.get("t") != "_localImage":
                    continue
                relation_index = int(rv.find("{*}v").text)
                image_path = targets[ids[relation_index]]
                if image_path in names:
                    found.append(ExcelPicture(sheet_name, cell.get("r", ""), archive.read(image_path)))
            except (AttributeError, IndexError, KeyError, TypeError, ValueError):
                continue
    return found


def visual_counts(path: Path) -> dict[str, int]:
    """Count actual pictures, charts and text-bearing shapes; ignore empty shapes."""
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        drawings = [name for name in names if name.startswith("xl/drawings/drawing")
                    and name.endswith(".xml")]
        pictures = 0
        text_shapes = 0
        for name in drawings:
            root = ET.fromstring(archive.read(name))
            pictures += len(root.findall(".//{*}pic"))
            text_shapes += sum(
                1 for shape in root.findall(".//{*}sp")
                if any((node.text or "").strip() for node in shape.findall(".//{*}t"))
            )
        rich = _rich_pictures(archive)
        # Rich relations without a readable cell mapping still represent unresolved images.
        rich_relations = 0
        rel_path = "xl/richData/_rels/richValueRel.xml.rels"
        if rel_path in names:
            rich_relations = sum(
                "image" in rel.get("Type", "")
                for rel in ET.fromstring(archive.read(rel_path))
            )
        return {
            "pictures": pictures + max(len(rich), rich_relations),
            "charts": sum(name.startswith("xl/charts/chart") and name.endswith(".xml") for name in names),
            "text_shapes": text_shapes,
        }


def rich_picture_cells(path: Path) -> set[tuple[str, str]]:
    with zipfile.ZipFile(path) as archive:
        return {(picture.sheet, picture.cell) for picture in _rich_pictures(archive)}


def extract_text_shapes(path: Path) -> list[dict]:
    """Read native shape text through sheet relations, retaining paragraph boundaries."""
    from openpyxl.utils import get_column_letter

    result = []
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        rels = {r.get('Id'): _part('xl/workbook.xml', r.get('Target', ''))
                for r in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        for sheet in ET.fromstring(archive.read('xl/workbook.xml')).findall('.//{*}sheet'):
            relation = next((v for k, v in sheet.attrib.items() if k.endswith('}id')), '')
            sheet_part = rels.get(relation, '')
            rel_part = posixpath.join(posixpath.dirname(sheet_part), '_rels', posixpath.basename(sheet_part) + '.rels')
            if rel_part not in names:
                continue
            for rel in ET.fromstring(archive.read(rel_part)):
                if not rel.get('Type', '').endswith('/drawing') or rel.get('TargetMode') == 'External':
                    continue
                drawing = _part(sheet_part, rel.get('Target', ''))
                if drawing not in names:
                    continue
                for anchor in ET.fromstring(archive.read(drawing)):
                    origin = anchor.find('{*}from')
                    cell = ''
                    if origin is not None:
                        cell = f"{get_column_letter(int(origin.findtext('{*}col', '0'))+1)}{int(origin.findtext('{*}row', '0'))+1}"
                    for shape in anchor.findall('.//{*}sp'):
                        paragraphs = [''.join(
                            (node.text or '') if node.tag.endswith('}t') else '\n'
                            for node in p.iter() if node.tag.endswith(('}t', '}br'))
                        ).strip() for p in shape.findall('.//{*}txBody/{*}p')]
                        text = '\n'.join(p for p in paragraphs if p)
                        if text:
                            result.append({'sheet': sheet.get('name', ''), 'cell': cell,
                                           'text': text, 'source_part': drawing})
    return result


def extract_pictures(path: Path) -> list[ExcelPicture]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    found = []
    with path.open("rb") as handle:
        book = load_workbook(handle, read_only=False)
        try:
            for sheet in book:
                for image in sheet._images:
                    anchor = image.anchor
                    cell = (
                        f"{get_column_letter(anchor._from.col + 1)}{anchor._from.row + 1}"
                        if hasattr(anchor, "_from") else str(anchor)
                    )
                    found.append(ExcelPicture(sheet.title, cell, image._data()))
        finally:
            book.close()
    with zipfile.ZipFile(path) as archive:
        found.extend(_rich_pictures(archive))
    return found
