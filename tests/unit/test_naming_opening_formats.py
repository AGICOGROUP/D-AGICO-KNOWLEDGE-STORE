from openpyxl import Workbook
from pptx import Presentation

from agico_kb.naming import _content_hint


def test_excel_opening_reads_extensionless_stored_original(tmp_path):
    source = tmp_path / 'original.xlsx'
    workbook = Workbook()
    workbook.active.append(['Greenfield Kiln', 'Equipment schedule'])
    workbook.save(source)
    stored = tmp_path / 'stored-blob'
    source.rename(stored)

    hint = _content_hint(stored, 'Annex-Equipment.xlsx')

    assert 'Greenfield Kiln' in hint
    assert 'Equipment schedule' in hint


def test_powerpoint_opening_reads_first_two_slides(tmp_path):
    source = tmp_path / 'stored-blob'
    deck = Presentation()
    for title in ['Kiln service guide', 'Maintenance schedule', 'Not opening text']:
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = title
    deck.save(source)

    hint = _content_hint(source, 'Guide.pptx')

    assert 'Kiln service guide' in hint
    assert 'Maintenance schedule' in hint
    assert 'Not opening text' not in hint


def test_large_pdf_still_reads_bounded_opening_pages(tmp_path):
    import pymupdf

    source = tmp_path / 'stored-blob'
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((50, 70), 'Cedar kiln design specification, installation and maintenance requirements')
        document.save(source)
    with source.open('ab') as stream:
        stream.write(b'\n%' + b' ' * (26 * 1024 * 1024))

    assert 'Cedar kiln design specification' in _content_hint(source, 'Specification.pdf')


def test_pdf_layout_whitespace_does_not_suppress_opening_ocr(tmp_path, monkeypatch):
    from types import SimpleNamespace

    import pymupdf

    source = tmp_path / 'stored-blob'
    source.write_bytes(b'synthetic fixture')
    page = SimpleNamespace(
        get_text=lambda **kwargs: 'M' + ' ' * 1800 + '\nM',
        get_pixmap=lambda **kwargs: SimpleNamespace(tobytes=lambda fmt: b'synthetic image'),
    )

    class Document:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def __len__(self):
            return 1

        def __getitem__(self, index):
            return page

    monkeypatch.setattr(pymupdf, 'open', lambda path: Document())
    from agico_kb import ingestion
    monkeypatch.setattr(ingestion, '_ocr_engine', lambda: lambda image:
                        SimpleNamespace(txts=['Cedar kiln process flow diagram']))

    assert 'Cedar kiln process flow diagram' in _content_hint(source, 'Flow.pdf')
