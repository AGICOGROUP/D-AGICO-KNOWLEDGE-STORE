from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorksheetImage
from PIL import Image, ImageDraw

from agico_kb.config import Settings
from agico_kb.excel_visuals import extract_pictures, visual_counts
from agico_kb.ingestion import parse_file
from agico_kb.vision import enrich_excel


def _workbook_with_image(path: Path, label: str):
    picture = path.with_suffix(".png")
    image = Image.new("RGB", (500, 220), "white")
    ImageDraw.Draw(image).text((30, 60), label, fill="black")
    image.save(picture)
    book = Workbook()
    sheet = book.active
    sheet.title = "参数表"
    sheet["A1"] = "型号"
    sheet["A2"] = "AX-210"
    sheet.add_image(WorksheetImage(picture), "C4")
    book.save(path)


def _settings(tmp_path, enabled=True):
    key = tmp_path / "api.txt"
    key.write_text("api=test-secret\n", encoding="utf-8")
    return Settings(
        "unused", tmp_path / "storage",
        vision_config_file=key if enabled else None,
    )


def test_excel_text_image_is_transcribed_and_located(tmp_path):
    path = tmp_path / "text-image.xlsx"
    _workbook_with_image(path, "AX-210 80 C")
    initial = parse_file(path, path.name)
    assert initial.status == "partial"
    assert visual_counts(path)["pictures"] == 1
    calls = []

    def fake(key, model, png, prompt):
        calls.append((key, model, png, prompt))
        return "图片写明 AX-210 允许工作温度为 80 C。"

    blob = tmp_path / "opaque.blob"
    blob.write_bytes(path.read_bytes())
    parsed = enrich_excel(
        initial, blob, _settings(tmp_path), call_model=fake,
        ocr_text=lambda _: "AX-210 80 C",
    )
    assert parsed.status == "ready"
    assert len(calls) == 1
    assert calls[0][1] == "glm-5.3-flash"
    visual = [c for c in parsed.chunks if c.locator["kind"] in {"image_ocr", "image_description"}]
    assert len(visual) == 2
    assert all(c.locator["sheet"] == "参数表" and c.locator["cell"] == "C4" for c in visual)
    assert "AX-210 80 C" in visual[0].text
    assert not any("未解释" in warning for warning in parsed.warnings)


def test_excel_logo_is_skipped_without_api_call(tmp_path):
    path = tmp_path / "logo.xlsx"
    _workbook_with_image(path, "AGICO CEMENT")
    initial = parse_file(path, path.name)

    def no_call(*args):
        raise AssertionError("A logo must not use the visual API")

    parsed = enrich_excel(
        initial, path, _settings(tmp_path), call_model=no_call,
        ocr_text=lambda _: "AGICO CEMENT",
    )
    assert parsed.status == "ready"
    assert not any(c.locator["kind"].startswith("image") for c in parsed.chunks)


def test_excel_empty_drawings_do_not_raise_missing_picture_warning(tmp_path):
    original = tmp_path / "plain.xlsx"
    path = tmp_path / "drawings.xlsx"
    book = Workbook()
    book.active["A1"] = "设备"
    book.save(original)
    with ZipFile(original) as source, ZipFile(path, "w") as target:
        for entry in source.infolist():
            target.writestr(entry, source.read(entry.filename))
        target.writestr(
            "xl/drawings/drawing1.xml",
            '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">'
            "<xdr:twoCellAnchor><xdr:sp/></xdr:twoCellAnchor></xdr:wsDr>",
        )
    assert visual_counts(path) == {"pictures": 0, "charts": 0, "text_shapes": 0}
    assert parse_file(path, path.name).status == "ready"


def test_rich_data_in_cell_picture_keeps_its_cell_location(tmp_path):
    original = tmp_path / "base.xlsx"
    path = tmp_path / "rich.xlsx"
    book = Workbook()
    book.active.title = "图片表"
    book.active["J2"] = "#VALUE!"
    book.save(original)
    image = Image.new("RGB", (300, 120), "white")
    ImageDraw.Draw(image).text((20, 40), "AX-210", fill="black")
    output = BytesIO()
    image.save(output, "PNG")
    with ZipFile(original) as source, ZipFile(path, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'r="J2"', b'r="J2" vm="1"')
            target.writestr(entry, data)
        target.writestr("xl/media/image1.png", output.getvalue())
        target.writestr(
            "xl/metadata.xml",
            '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:xlrd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata">'
            '<futureMetadata><bk><extLst><ext><xlrd:rvb i="0"/></ext></extLst></bk></futureMetadata>'
            '<valueMetadata><bk><rc v="0"/></bk></valueMetadata></metadata>',
        )
        target.writestr(
            "xl/richData/rdrichvalue.xml",
            '<rvData xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata">'
            '<rv s="0"><v>0</v></rv></rvData>',
        )
        target.writestr(
            "xl/richData/rdrichvaluestructure.xml",
            '<rvStructures xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata">'
            '<s t="_localImage"/></rvStructures>',
        )
        target.writestr(
            "xl/richData/richValueRel.xml",
            '<richValueRels xmlns="http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<rel r:id="rId1"/></richValueRels>',
        )
        target.writestr(
            "xl/richData/_rels/richValueRel.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="image" Target="../media/image1.png"/>'
            '</Relationships>',
        )
    pictures = extract_pictures(path)
    assert len(pictures) == 1
    assert (pictures[0].sheet, pictures[0].cell) == ("图片表", "J2")
    assert visual_counts(path)["pictures"] == 1
    parsed = parse_file(path, path.name)
    assert "#VALUE!" not in "\n".join(chunk.text for chunk in parsed.chunks)
