import hashlib
import zipfile

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from agico_kb.ingestion import parse_file


@pytest.mark.parametrize("dimension", ["inflated", "missing", "understated"])
def test_actual_content_survives_false_dimensions_and_style_only_cells(tmp_path, dimension):
    original = tmp_path / "original.xlsx"
    path = tmp_path / "uploaded.blob"
    book = Workbook()
    sheet = book.active
    sheet["B3"] = "设备"
    sheet["D3"] = "数量"
    sheet["B877"] = "风机"
    sheet["D877"] = 0
    sheet["AR877"] = "=D877*3"
    sheet["XFD1048576"].fill = PatternFill("solid", fgColor="FF0000")
    book.save(original)
    import re
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(path, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                replacement = b"" if dimension == "missing" else b'<dimension ref="A1"/>' if dimension == "understated" else b'<dimension ref="A1:XFD1048576"/>'
                data = re.sub(rb"<dimension\b[^>]*/>", replacement, data)
            target.writestr(entry, data)
    digest = hashlib.sha256(path.read_bytes()).digest()
    result = parse_file(path, "fresh.xlsx")
    assert hashlib.sha256(path.read_bytes()).digest() == digest
    assert len(result.chunks) == 1
    assert result.chunks[0].locator["range"] == "B3:AR877"
    assert result.chunks[0].locator["columns"] == ["B", "D", "AR"]
    assert result.chunks[0].locator["source_rows"] == [3, 877]
    assert all(v in result.chunks[0].text for v in ["设备", "风机", "0", "=D877*3", "缓存结果缺失"])
    assert "XFD" not in result.chunks[0].text


def test_style_only_sheet_produces_no_separator_chunks(tmp_path):
    book = Workbook()
    book.active["XFD1048576"].fill = PatternFill("solid", fgColor="FF0000")
    path = tmp_path / "blank.xlsx"
    book.save(path)
    assert not parse_file(path, path.name).chunks


def test_true_large_content_extent_remains_bounded(tmp_path):
    book = Workbook()
    book.active.cell(1001, 1000, "real data")
    path = tmp_path / "large.xlsx"
    book.save(path)
    with pytest.raises(ValueError, match="保护上限"):
        parse_file(path, path.name)
