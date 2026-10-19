import hashlib
import tempfile
import time

import pymupdf
import pytest
from PIL import Image

from agico_kb.errors import KBError
from agico_kb.previews import DocumentPreviewService, prepare_document, render_page
from agico_kb.thumbnails import _run, office_converter


def fresh_document(folder, suffix):
    path = folder / ('fresh' + suffix)
    if suffix == '.pdf':
        with pymupdf.open() as doc:
            for index in range(3):
                page = doc.new_page(width=300, height=400)
                page.insert_text((30, 40), f'PAGE {index + 1} END-OF-DOCUMENT' if index == 2 else f'PAGE {index + 1}')
            doc.save(path)
    elif suffix == '.docx':
        from docx import Document
        doc = Document()
        doc.add_paragraph('FIRST PAGE')
        doc.add_page_break()
        doc.add_paragraph('LAST PAGE END-OF-DOCUMENT')
        doc.save(path)
    elif suffix == '.xlsx':
        from openpyxl import Workbook
        doc = Workbook()
        doc.active['A1'] = 'FIRST SHEET'
        doc.create_sheet('Second')['A1'] = 'LAST SHEET END-OF-DOCUMENT'
        doc.save(path)
    elif suffix == '.pptx':
        from pptx import Presentation
        doc = Presentation()
        for text in ['FIRST SLIDE', 'LAST SLIDE END-OF-DOCUMENT']:
            slide = doc.slides.add_slide(doc.slide_layouts[5])
            slide.shapes.title.text = text
        doc.save(path)
    else:
        images = [Image.new('RGB', (900, 600), color) for color in ['red', 'blue']]
        images[0].save(path, save_all=True, append_images=images[1:])
    return path


@pytest.mark.parametrize('suffix', ['.pdf', '.docx', '.xlsx', '.pptx', '.tiff'])
def test_full_preview_includes_last_page_and_preserves_original(tmp_path, suffix):
    if suffix in {'.docx', '.xlsx', '.pptx'} and not office_converter():
        pytest.skip('LibreOffice not installed')
    source = fresh_document(tmp_path, suffix)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    target = tmp_path / 'derived'
    manifest = prepare_document(source, source.name, target)
    expected = 3 if suffix == '.pdf' else 2
    assert manifest['pages'] == expected
    assert (target / 'manifest.json').is_file()
    if suffix in {'.docx', '.xlsx', '.pptx'}:
        with pymupdf.open(target / 'document.pdf') as doc:
            assert 'END-OF-DOCUMENT' in doc[-1].get_text()
    image = target / 'last.png'
    render_page(source, source.name, target, expected, 1800, image)
    with Image.open(image) as rendered:
        assert max(rendered.size) > 512
        assert max(rendered.size) <= 1800
        if suffix == '.tiff':
            assert rendered.getpixel((100, 100))[:3] == (0, 0, 255)
    with pytest.raises(ValueError):
        render_page(source, source.name, target, expected + 1, 1800, image)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest


def test_background_preparation_and_page_cache_survive_service_restart(tmp_path):
    source = fresh_document(tmp_path, '.pdf')
    service = DocumentPreviewService(tmp_path / 'storage')
    try:
        def ready(callback):
            result = callback()
            deadline = time.monotonic() + 20
            while result is None and time.monotonic() < deadline:
                time.sleep(.05)
                result = callback()
            assert result is not None
            return result
        assert ready(lambda: service.manifest(source, source.name))['pages'] == 3
        page = ready(lambda: service.page(source, source.name, 3, 1800))
        assert page.is_file()
        with pytest.raises(KBError) as error:
            service.page(source, source.name, 4, 1800)
        assert error.value.status == 404
    finally:
        service.close()
    second = DocumentPreviewService(tmp_path / 'storage')
    try:
        assert second.manifest(source, source.name)['pages'] == 3
        assert second.page(source, source.name, 3, 1800) == page
    finally:
        second.close()


def test_unreadable_preview_reports_error_instead_of_empty_success(tmp_path):
    source = tmp_path / 'bad.pdf'
    source.write_bytes(b'broken')
    service = DocumentPreviewService(tmp_path / 'cache')
    try:
        assert service.manifest(source, source.name) is None
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            try:
                service.manifest(source, source.name)
            except KBError as error:
                assert error.status == 422
                assert 'bad.pdf' not in str(error.payload)
                break
            time.sleep(.05)
        else:
            pytest.fail('Renderer failure was not reported')
    finally:
        service.close()


def test_excel_full_preview_does_not_omit_cells_outside_print_area(tmp_path):
    if not office_converter():
        pytest.skip('LibreOffice not installed')
    from openpyxl import Workbook
    source = tmp_path / 'print-area.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = 'PRINT AREA'
    workbook.active.print_area = 'A1'
    workbook.active['A20'] = 'OUTSIDE-PRINT-AREA-REQUIRED'
    workbook.active.column_dimensions['A'].width = 50
    workbook.save(source)
    before = source.read_bytes()
    prepare_document(source, source.name, tmp_path / 'derived')
    with pymupdf.open(tmp_path / 'derived/document.pdf') as doc:
        assert 'OUTSIDE-PRINT-AREA-REQUIRED' in ''.join(page.get_text() for page in doc)
    assert source.read_bytes() == before


@pytest.mark.parametrize('original,legacy,export', [
    ('.docx', '.doc', 'doc:MS Word 97'),
    ('.xlsx', '.xls', 'xls:MS Excel 97'),
    ('.pptx', '.ppt', 'ppt:MS PowerPoint 97'),
])
def test_legacy_office_full_preview(original, legacy, export, tmp_path):
    converter = office_converter()
    if not converter:
        pytest.skip('LibreOffice not installed')
    source = fresh_document(tmp_path, original)
    old = tmp_path / 'legacy'
    old.mkdir()
    with tempfile.TemporaryDirectory() as profile:
        from pathlib import Path
        _run([converter, f'-env:UserInstallation={Path(profile).as_uri()}', '--headless',
              '--convert-to', export, '--outdir', str(old), str(source)])
    source = old / ('fresh' + legacy)
    before = source.read_bytes()
    folder = tmp_path / 'derived'
    assert prepare_document(source, source.name, folder)['pages'] == 2
    with pymupdf.open(folder / 'document.pdf') as doc:
        assert 'END-OF-DOCUMENT' in doc[-1].get_text()
    assert source.read_bytes() == before
