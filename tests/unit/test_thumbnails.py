import shutil
import subprocess
import threading
from io import BytesIO

import pytest
from PIL import Image

from agico_kb.errors import KBError
from agico_kb.thumbnails import (
    ThumbnailService,
    blender_converter,
    generate_thumbnail,
    image_thumbnail,
    office_converter,
)


def test_blend_glb_and_step_preview_show_new_model_without_changing_source(tmp_path):
    blender = blender_converter()
    if not blender:
        pytest.skip("Blender not installed")
    script = tmp_path / "make_test_models.py"
    script.write_text(
        "import bpy, sys\n"
        "from pathlib import Path\n"
        "folder = Path(sys.argv[sys.argv.index('--') + 1])\n"
        "bpy.ops.mesh.primitive_cube_add(size=2)\n"
        "bpy.ops.wm.save_as_mainfile(filepath=str(folder / 'fresh.blend'))\n"
        "bpy.ops.export_scene.gltf(filepath=str(folder / 'fresh.glb'), export_format='GLB')\n",
        encoding="utf-8",
    )
    subprocess.run([blender, "--background", "--factory-startup", "--disable-autoexec",
                    "--python-exit-code", "1", "--python", str(script), "--", str(tmp_path)],
                   check=True, timeout=30, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    from OCP.BRepPrimAPI import BRepPrimAPI_MakeBox
    from OCP.STEPControl import STEPControl_AsIs, STEPControl_Writer

    writer = STEPControl_Writer()
    writer.Transfer(BRepPrimAPI_MakeBox(10, 8, 6).Shape(), STEPControl_AsIs)
    writer.Write(str(tmp_path / "fresh.step"))

    for suffix in [".blend", ".glb", ".step"]:
        source = tmp_path / ("fresh" + suffix)
        before = source.read_bytes()
        target = tmp_path / ("preview-" + suffix[1:] + ".png")
        generate_thumbnail(source, source.name, target)
        assert source.read_bytes() == before
        with Image.open(target) as image:
            assert image.size == (512, 512)
            assert image.convert("L").getextrema()[0] < 100


def test_thumbnail_keeps_aspect_transparency_and_original(tmp_path):
    path = tmp_path / "original.blob"
    Image.new("RGBA", (1200, 600), (255, 255, 255, 0)).save(path, format="PNG")
    before = path.read_bytes()
    with Image.open(BytesIO(image_thumbnail(path, "白色LOGO.PNG"))) as preview:
        assert preview.size == (512, 256)
        assert preview.getpixel((0, 0))[3] == 0
    assert path.read_bytes() == before


def test_thumbnail_handles_unreadable_and_non_image_files(tmp_path):
    path = tmp_path / "original.blob"
    path.write_bytes(b"not an image")
    for name in ["broken.png", "document.pdf"]:
        with pytest.raises(KBError) as error:
            image_thumbnail(path, name)
        assert error.value.status == 422


def test_dxf_preview_prefers_nonempty_layout_and_preserves_source(tmp_path):
    import ezdxf

    source = tmp_path / "new.dxf"
    doc = ezdxf.new()
    doc.modelspace().add_circle((0, 0), 50)
    paper = doc.layouts.new("Sheet")
    paper.add_lwpolyline([(0, 0), (100, 0), (100, 60), (0, 60)], close=True)
    paper.add_text("DRAWING A-01", dxfattribs={"height": 6, "insert": (5, 25)})
    doc.saveas(source)
    before = source.read_bytes()
    first = tmp_path / "first.png"
    generate_thumbnail(source, "new.dxf", first)
    assert source.read_bytes() == before
    with Image.open(first) as image:
        assert max(image.size) <= 512
        assert image.convert("L").getextrema()[0] < 100
    # Rendering just that layout independently yields the same preview.
    doc.modelspace().delete_all_entities()
    only_sheet = tmp_path / "only-sheet.dxf"
    doc.saveas(only_sheet)
    second = tmp_path / "second.png"
    generate_thumbnail(only_sheet, "only-sheet.dxf", second)
    assert first.read_bytes() == second.read_bytes()


def test_video_preview_uses_real_frame_and_keeps_original(tmp_path):
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        pytest.skip("FFmpeg not installed")
    source, preview = tmp_path / "fresh.mp4", tmp_path / "preview.png"
    subprocess.run(
        [
            ffmpeg,
            "-nostdin",
            "-v",
            "error",
            "-f",
            "lavfi",
            "-i",
            "color=c=red:s=320x180:d=1",
            "-c:v",
            "mpeg4",
            "-y",
            str(source),
        ],
        check=True,
        timeout=15,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    before = source.read_bytes()
    generate_thumbnail(source, "fresh.mp4", preview)
    with Image.open(preview) as image:
        assert image.width / image.height == pytest.approx(16 / 9, abs=0.02)
        red, green, blue = image.convert("RGB").getpixel((image.width // 2, image.height // 2))
        assert red > 200 and green < 30 and blue < 30
    assert source.read_bytes() == before


def test_cad_inactive_paper_viewport_falls_back_to_model(tmp_path):
    import ezdxf

    source, preview = tmp_path / "inactive.dxf", tmp_path / "preview.png"
    doc = ezdxf.new()
    doc.modelspace().add_circle((0, 0), 10)
    doc.layouts.get("Layout1").add_viewport(
        center=(50, 50), size=(100, 100), view_center_point=(0, 0), view_height=100, status=0
    )
    doc.saveas(source)
    generate_thumbnail(source, "inactive.dxf", preview)
    with Image.open(preview) as image:
        assert image.convert("L").getextrema()[0] < 100


def test_background_preview_deduplicates_and_reuses_persistent_cache(tmp_path):
    source = tmp_path / "source.blob"
    Image.new("RGB", (24, 12), "red").save(source, format="PNG")
    service = ThumbnailService(tmp_path)
    try:
        assert service.get(source, "new.png") is None
        key, job = next(iter(service.jobs.items()))
        service.get(source, "new.png")
        assert len(service.jobs) <= 1
        assert job.result(timeout=10)
        cached = service.get(source, "new.png")
        assert cached.is_file()
        assert cached.name == key + ".png"
        restarted = ThumbnailService(tmp_path)
        try:
            assert restarted.get(source, "new.png") == cached
            assert not restarted.jobs
        finally:
            restarted.close()
    finally:
        service.close()


def test_background_decoder_failure_returns_fallback_error(tmp_path):
    source = tmp_path / "broken.blob"
    source.write_bytes(b"bad file")
    service = ThumbnailService(tmp_path)
    try:
        assert service.get(source, "broken.mp4") is None
        job = next(iter(service.jobs.values()))
        assert not job.result(timeout=10)
        with pytest.raises(KBError) as error:
            service.get(source, "broken.mp4")
        assert error.value.status == 422
    finally:
        service.close()


def test_pdf_preview_is_first_page_and_keeps_original(tmp_path):
    import pymupdf

    source, preview = tmp_path / "fresh.pdf", tmp_path / "preview.png"
    doc = pymupdf.open()
    for color in [(1, 0, 0), (0, 0, 1)]:
        page = doc.new_page(width=600, height=800)
        page.draw_rect(page.rect, color=color, fill=color)
    doc.save(source)
    doc.close()
    before = source.read_bytes()
    generate_thumbnail(source, "fresh.pdf", preview)
    with Image.open(preview) as image:
        assert image.size == (384, 512)
        red, green, blue = image.convert("RGB").getpixel((100, 100))
        assert red > 200 and blue < 30 and green < 30
    assert source.read_bytes() == before


@pytest.mark.parametrize("suffix", [".docx", ".xlsx", ".pptx"])
def test_office_first_page_preview_keeps_original(tmp_path, suffix):
    if not office_converter():
        pytest.skip("LibreOffice not installed")
    source, preview = tmp_path / ("fresh" + suffix), tmp_path / "preview.png"
    if suffix == ".docx":
        from docx import Document

        doc = Document()
        doc.add_heading("Fresh document preview", 0)
        doc.add_paragraph("首页预览验证：原文保留，型号 AX-100。")
        doc.add_page_break()
        doc.add_paragraph("Second page")
        doc.save(source)
    elif suffix == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill

        doc = Workbook()
        doc.active.append(["Product", "Quantity", "Price"])
        doc.active.append(["AX-100", 12, 350])
        for cell in doc.active[1]:
            cell.fill = PatternFill("solid", fgColor="0055AA")
        doc.save(source)
    else:
        from pptx import Presentation

        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[0])
        slide.shapes.title.text = "Fresh slide preview"
        slide.placeholders[1].text = "First slide, AX-100"
        doc.save(source)
    before = source.read_bytes()
    generate_thumbnail(source, source.name, preview)
    with Image.open(preview) as image:
        assert max(image.size) <= 512
        assert image.convert("L").getextrema()[0] < 100
    assert source.read_bytes() == before


def test_text_preview_renders_original_excerpt(tmp_path):
    source, preview = tmp_path / "fresh.txt", tmp_path / "preview.png"
    source.write_text("企业知识测试\n产品 AX-100，数量 12。\n" * 50, encoding="utf-8")
    before = source.read_bytes()
    generate_thumbnail(source, "fresh.txt", preview)
    with Image.open(preview) as image:
        assert image.convert("L").getextrema()[0] < 100
    assert source.read_bytes() == before


def test_encrypted_pdf_is_not_rendered_without_password(tmp_path):
    import pymupdf

    source = tmp_path / "locked.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(source, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="test")
    doc.close()
    with pytest.raises(ValueError):
        generate_thumbnail(source, source.name, tmp_path / "preview.png")


def test_pdf_preview_does_not_wait_for_slow_cad_conversions(tmp_path, monkeypatch):
    import pymupdf

    service = ThumbnailService(tmp_path)
    release = threading.Event()
    original_generate = service._generate

    def blocked_cad(path, name, target):
        if name.endswith(".dwg"):
            release.wait(10)
            return False
        return original_generate(path, name, target)

    monkeypatch.setattr(service, "_generate", blocked_cad)
    try:
        for index in range(2):
            source = tmp_path / f"slow-{index}.dwg"
            source.write_bytes(b"synthetic blocked conversion")
            assert service.get(source, source.name) is None
        pdf = tmp_path / "new.pdf"
        doc = pymupdf.open()
        doc.new_page().insert_text((30, 40), "Fast PDF")
        doc.save(pdf)
        doc.close()
        assert service.get(pdf, pdf.name) is None
        pdf_job = list(service.jobs.values())[-1]
        assert pdf_job.result(timeout=5)
        assert not release.is_set()
    finally:
        release.set()
        service.close()
