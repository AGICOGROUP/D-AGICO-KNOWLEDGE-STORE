import hashlib
from io import BytesIO

import pymupdf
import pytest
from docx import Document
from docx.shared import Cm
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from agico_kb.config import Settings
from agico_kb.ingestion import parse_file
from agico_kb.vision import (
    _has_image_text,
    _pdf_text_image_render,
    enrich_excel,
    enrich_image,
    enrich_pdf,
    enrich_ppt,
    enrich_word,
)


def picture(text=False):
    image = Image.new('RGB', (600, 400), '#719db5')
    if text:
        ImageDraw.Draw(image).text((40, 150), 'MODEL AX-210', fill='black',
                                  font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 42))
    stream = BytesIO()
    image.save(stream, format='PNG')
    return stream.getvalue()


def settings(tmp_path):
    config = tmp_path / 'vision.txt'
    config.write_text('api=test-secret\n', encoding='utf-8')
    return Settings('test', tmp_path, vision_config_file=config)


def forbidden_api(*_):
    pytest.fail('A no-text picture must never invoke the API')


def test_real_local_screen_distinguishes_label_and_no_text():
    assert _has_image_text(picture(True))
    assert not _has_image_text(picture(False))


def test_word_skips_no_text_figure_and_clears_only_figure_gap(tmp_path):
    path = tmp_path / 'photo.docx'
    doc = Document()
    doc.add_paragraph('Original body remains available.')
    doc.add_picture(BytesIO(picture()), width=Cm(12))
    doc.save(path)
    parsed = enrich_word(parse_file(path, path.name), path, settings(tmp_path), call_model=forbidden_api)
    assert not any(c.locator['kind'] == 'figure' for c in parsed.chunks)
    assert not any('未解释' in w for w in parsed.warnings)


def test_ppt_title_outside_photo_does_not_trigger_api(tmp_path):
    path = tmp_path / 'photo.pptx'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(.2), Inches(5), Inches(.4)).text = 'Readable slide heading'
    slide.shapes.add_picture(BytesIO(picture()), Inches(1), Inches(1), width=Inches(6))
    deck.save(path)
    parsed = enrich_ppt(parse_file(path, path.name), path, settings(tmp_path),
                        call_model=forbidden_api, render_slides=forbidden_api)
    assert any('Readable slide heading' in c.text for c in parsed.chunks)
    assert not any(c.locator['kind'] == 'slide_visual' for c in parsed.chunks)
    assert not any('未解释' in w for w in parsed.warnings)


def test_pdf_native_caption_outside_photo_does_not_trigger_api(tmp_path):
    path = tmp_path / 'photo.pdf'
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(40, 40, 550, 400), 'Readable native body. ' * 35)
        page.insert_image(pymupdf.Rect(40, 450, 550, 750), stream=picture())
        doc.save(path)
    parsed = enrich_pdf(parse_file(path, path.name), path, settings(tmp_path), call_model=forbidden_api)
    assert not any(c.locator['kind'] == 'page_visual' for c in parsed.chunks)
    assert not any('未解释' in w for w in parsed.warnings)
    assert not any('AI 转述' in w for w in parsed.warnings)


def test_standalone_no_text_image_keeps_original_without_summary(tmp_path):
    path = tmp_path / 'photo.png'
    path.write_bytes(picture())
    parsed = enrich_image(parse_file(path, path.name), path, settings(tmp_path), call_model=forbidden_api)
    assert parsed.status == 'stored_only'
    assert not parsed.chunks


def test_excel_no_text_picture_does_not_trigger_api(tmp_path):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as SheetImage

    path = tmp_path / 'photo.xlsx'
    book = Workbook()
    book.active['A1'] = 'Native data'
    book.active.add_image(SheetImage(BytesIO(picture())), 'C4')
    book.save(path)
    parsed = enrich_excel(parse_file(path, path.name), path, settings(tmp_path),
                          call_model=forbidden_api, ocr_text=lambda _: '')
    assert not any(c.locator['kind'].startswith('image') for c in parsed.chunks)
    assert not any('未解释' in w for w in parsed.warnings)


def test_ppt_mixed_pictures_filters_render_copy_only(tmp_path):
    path = tmp_path / 'mixed.pptx'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(BytesIO(picture(True)), Inches(.5), Inches(1), width=Inches(4))
    slide.shapes.add_picture(BytesIO(picture(False)), Inches(5), Inches(1), width=Inches(4))
    deck.save(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()

    def render(filtered, indices):
        assert filtered != path
        shapes = Presentation(filtered).slides[0].shapes
        assert len(shapes) == 1
        assert _has_image_text(shapes[0].image.blob)
        return {1: picture(True)}

    parsed = enrich_ppt(parse_file(path, path.name), path, settings(tmp_path),
                        render_slides=render, call_model=lambda *_: 'MODEL AX-210')
    assert any(c.locator['kind'] == 'slide_visual' for c in parsed.chunks)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_pdf_mixed_pictures_masks_wordless_region_only(tmp_path):
    path = tmp_path / 'mixed.pdf'
    with pymupdf.open() as doc:
        page = doc.new_page(width=600, height=800)
        page.insert_image(pymupdf.Rect(20, 50, 580, 350), stream=picture(True))
        page.insert_image(pymupdf.Rect(20, 450, 580, 750), stream=picture(False))
        doc.save(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    calls = []
    from agico_kb.ingestion import _content_images

    with pymupdf.open(path) as doc:
        blob, readable = _pdf_text_image_render(doc[0], 1, _content_images(doc[0]))
        assert readable
        assert Image.open(BytesIO(blob)).convert('RGB').getpixel((300, 600)) == (255, 255, 255)

    def model(_key, _model, blob, _prompt):
        image = Image.open(BytesIO(blob)).convert('RGB')
        # API normalization trims white margins, leaving only the top text picture.
        assert image.height < image.width
        assert _has_image_text(blob)
        calls.append(1)
        return 'MODEL AX-210'

    parsed = enrich_pdf(parse_file(path, path.name), path, settings(tmp_path), call_model=model)
    assert calls == [1, 1]  # transcription plus source-image review
    assert any(c.locator['kind'] == 'page_visual' for c in parsed.chunks)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
