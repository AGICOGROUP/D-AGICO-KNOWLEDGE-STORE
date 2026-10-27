import sqlite3
from io import BytesIO

import httpx
import pytest
from docx import Document
from docx.shared import Cm
from PIL import Image

from agico_kb.config import Settings
from agico_kb.ingestion import parse_file
from agico_kb.vision import _credentials, enrich_image, enrich_pdf, enrich_ppt, enrich_word


@pytest.fixture(autouse=True)
def readable_test_pictures(monkeypatch):
    # These tests isolate transport/cache/budget behavior using synthetic picture bytes.
    # Real text screening and no-text skips have their own source-based tests.
    monkeypatch.setattr("agico_kb.vision._has_image_text", lambda _: True)


def test_glm_is_default_for_key_only_config(tmp_path):
    config = tmp_path / "vision-api.txt"
    config.write_text("api=test-secret\n", encoding="utf-8")
    assert _credentials(config) == ("test-secret", "glm-5.3-flash")
    config.write_text("api=test-secret\nmodel=deepseek-v4.1-flash\n", encoding="utf-8")
    assert _credentials(config) == ("test-secret", "deepseek-v4.1-flash")


def _word_with_figures(path, count=1, color="black"):
    picture = path.with_suffix(".png")
    image = Image.new("RGB", (120, 80), "white")
    image.paste(color, (10, 10, 100, 60))
    image.save(picture)
    document = Document()
    document.add_paragraph("AX-210 安装说明")
    document.add_picture(str(picture), width=Cm(1.5))  # decorative logo
    for _ in range(count):
        document.add_picture(str(picture), width=Cm(12))
    document.save(path)


def _settings(tmp_path, *, limit=2, daily=20, enabled=True):
    config = tmp_path / "vision-api.txt"
    config.write_text(
        "api=test-secret\nmodel=glm-5.3-flash/deepseek-v4.1-flash/\n", encoding="utf-8"
    )
    return Settings(
        "unused", tmp_path / "storage",
        vision_config_file=config if enabled else None,
        vision_max_images_per_file=limit,
        vision_max_calls_per_day=daily,
    )


def test_word_vision_enriches_only_content_figures_and_uses_cache(tmp_path):
    path = tmp_path / "figures.docx"
    _word_with_figures(path, count=2)
    settings = _settings(tmp_path)
    calls = []

    def fake(key, model, png):
        calls.append((key, model, png))
        assert Image.open(BytesIO(png)).size[0] >= 100
        return "图中标有 AX-210 和接线端子。"

    parsed = enrich_word(parse_file(path, path.name), path, settings, call_model=fake)
    assert parsed.status == "ready"
    assert len([c for c in parsed.chunks if c.locator["kind"] == "figure"]) == 2
    assert len(calls) == 1  # same image is analyzed once, even at two positions
    assert not any("未解释" in w for w in parsed.warnings)

    # The worker stores original documents under opaque .blob names.
    blob_path = tmp_path / "opaque.blob"
    blob_path.write_bytes(path.read_bytes())
    via_blob = enrich_word(parse_file(path, path.name), blob_path, settings, call_model=fake)
    assert via_blob.status == "ready"

    again = enrich_word(parse_file(path, path.name), path, settings, call_model=fake)
    assert again.status == "ready"
    assert len(calls) == 1  # reindexing also reads the persistent cache


def test_word_floating_content_picture_is_also_described(tmp_path):
    from docx.oxml.ns import qn

    picture = tmp_path / "diagram.png"
    Image.new("RGB", (600, 400), "white").save(picture)
    document = Document()
    document.add_paragraph("设备布局")
    shape = document.add_picture(str(picture), width=Cm(12))
    shape._inline.tag = qn("wp:anchor")
    path = tmp_path / "floating.docx"
    document.save(path)
    parsed = parse_file(path, path.name)
    assert any("图片" in warning and "未解释" in warning for warning in parsed.warnings)
    parsed = enrich_word(
        parsed, path, _settings(tmp_path), call_model=lambda *_: "图中标出设备主电机。"
    )
    assert parsed.status == "ready"
    assert any("主电机" in chunk.text for chunk in parsed.chunks)


def test_word_vision_disabled_or_limited_preserves_partial_warning(tmp_path):
    path = tmp_path / "limited.docx"
    _word_with_figures(path, count=2)
    called = []

    def fake(*args):
        called.append(1)
        return "图形摘要"

    disabled = enrich_word(
        parse_file(path, path.name), path,
        _settings(tmp_path, enabled=False), call_model=fake,
    )
    assert disabled.status == "partial"
    assert called == []

    limited = enrich_word(
        parse_file(path, path.name), path,
        _settings(tmp_path, limit=1), call_model=fake,
    )
    assert limited.status == "partial"
    assert len(called) == 1
    assert any("1 张图未解释" in w for w in limited.warnings)


def test_word_vision_budget_and_api_failure_do_not_lose_text(tmp_path):
    path = tmp_path / "budget.docx"
    _word_with_figures(path)
    called = []

    def fake(*args):
        called.append(1)
        raise RuntimeError("API failure with secret")

    from agico_kb.vision import _reserve, _store
    cfg = _settings(tmp_path, daily=1)
    conn = _store(cfg.storage_root)
    assert _reserve(conn, 1)
    conn.close()
    blocked = enrich_word(
        parse_file(path, path.name), path,
        cfg, call_model=fake,
    )
    assert called == []
    assert blocked.status == "partial"
    assert any("AX-210" in c.text for c in blocked.chunks)

    failed = enrich_word(
        parse_file(path, path.name), path,
        _settings(tmp_path), call_model=fake,
    )
    assert len(called) == 1
    assert failed.status == "partial"
    assert all("secret" not in w for w in failed.warnings)


def test_transient_visual_failure_recovers_in_the_same_ingestion(tmp_path, monkeypatch):
    path = tmp_path / "retry.docx"
    _word_with_figures(path)
    settings = _settings(tmp_path)
    monkeypatch.setattr("agico_kb.vision.time.sleep", lambda _: None)
    attempts = []

    def fake(*_args):
        attempts.append(1)
        if len(attempts) == 1:
            raise httpx.ReadTimeout("temporary timeout")
        return "图中标有 AX-210。"

    parsed = enrich_word(parse_file(path, path.name), path, settings, call_model=fake)
    assert parsed.status == "ready"
    assert len(attempts) == 2
    assert any("AX-210" in chunk.text for chunk in parsed.chunks)
    with sqlite3.connect(settings.storage_root / ".vision" / "cache.sqlite3") as conn:
        assert conn.execute("SELECT count FROM daily_calls").fetchone()[0] == 2
        assert conn.execute("SELECT count(*) FROM summaries").fetchone()[0] == 1


def test_transient_visual_failure_respects_daily_budget(tmp_path, monkeypatch):
    path = tmp_path / "budget-retry.docx"
    _word_with_figures(path)
    monkeypatch.setattr("agico_kb.vision.time.sleep", lambda _: None)
    attempts = []

    def fake(*_args):
        attempts.append(1)
        raise httpx.ReadTimeout("temporary timeout")

    parsed = enrich_word(
        parse_file(path, path.name), path, _settings(tmp_path, daily=1), call_model=fake
    )
    assert len(attempts) == 1
    assert parsed.status == "partial"
    assert any("未解释" in warning for warning in parsed.warnings)


def test_daily_budget_is_shared_between_files(tmp_path):
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    _word_with_figures(first, color="black")
    _word_with_figures(second, color="blue")
    settings = _settings(tmp_path, daily=1)
    calls = []

    def fake(*args):
        calls.append(1)
        return "可检索的图形摘要"

    one = enrich_word(parse_file(first, first.name), first, settings, call_model=fake)
    two = enrich_word(parse_file(second, second.name), second, settings, call_model=fake)
    assert one.status == "ready"
    assert two.status == "partial"
    assert len(calls) == 1


def test_explicit_bulk_limit_can_cover_twenty_four_figures(tmp_path):
    path = tmp_path / "large.docx"
    _word_with_figures(path, count=24)
    calls = []

    def fake(*args):
        calls.append(1)
        return "图形摘要"

    parsed = enrich_word(
        parse_file(path, path.name), path, _settings(tmp_path, limit=32), call_model=fake
    )
    assert parsed.status == "ready"
    assert sum(c.locator["kind"] == "figure" for c in parsed.chunks) == 24
    assert len(calls) == 1


def test_ppt_picture_gets_searchable_visual_text_and_decorations_do_not_warn(tmp_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    picture = tmp_path / "diagram.png"
    Image.new("RGB", (640, 400), "white").save(picture)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(1), Inches(.1))
    slide.shapes.add_picture(str(picture), Inches(1), Inches(1), width=Inches(6))
    path = tmp_path / "diagram.pptx"
    deck.save(path)
    parsed = enrich_ppt(
        parse_file(path, path.name), path, _settings(tmp_path),
        call_model=lambda *_: "图中标有 AX-210 接线端子。",
    )
    assert parsed.status == "ready"
    assert any("AX-210" in c.text and c.locator["slide"] == 1 for c in parsed.chunks)
    assert not any("未解释" in warning for warning in parsed.warnings)


def test_ppt_visual_inventory_covers_slides_beyond_warning_display_limit(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    picture = tmp_path / "diagram.png"
    Image.new("RGB", (400, 300), "white").save(picture)
    deck = Presentation()
    for _ in range(23):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(str(picture), Inches(1), Inches(1), width=Inches(6))
    path = tmp_path / "long.pptx"
    deck.save(path)
    parsed = parse_file(path, path.name)
    assert sum('含未解释' in warning for warning in parsed.warnings) == 23
    parsed = enrich_ppt(
        parsed, path, _settings(tmp_path, limit=32),
        call_model=lambda *_: "设备布局示意图",
        render_slides=lambda _path, indices: {index: picture.read_bytes() for index in indices},
    )
    assert sum(c.locator["kind"] == "slide_visual" for c in parsed.chunks) == 23
    assert parsed.status == "ready"
    assert not any("未展开" in warning or "未解释" in warning for warning in parsed.warnings)


def test_visual_only_ppt_keeps_partial_index_when_daily_limit_stops_early(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    picture = tmp_path / "diagram.png"
    Image.new("RGB", (400, 300), "white").save(picture)
    deck = Presentation()
    for _ in range(2):
        deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(
            str(picture), Inches(1), Inches(1), width=Inches(6)
        )
    path = tmp_path / "visual-only.pptx"
    deck.save(path)
    parsed = enrich_ppt(
        parse_file(path, path.name), path, _settings(tmp_path, limit=1),
        call_model=lambda *_: "图形信息",
        render_slides=lambda _path, indices: {index: picture.read_bytes() for index in indices},
    )
    assert parsed.status == "partial"
    assert len(parsed.chunks) == 1
    assert any("1 页视觉内容未解释" in warning for warning in parsed.warnings)


def test_standalone_image_keeps_ocr_and_adds_model_description(tmp_path):
    from agico_kb.ingestion import Chunk, Parsed
    picture = tmp_path / "layout.png"
    Image.new("RGB", (800, 500), "white").save(picture)
    parsed = enrich_image(
        Parsed([Chunk("AX-210", {"kind": "ocr", "image": 1})]), picture, _settings(tmp_path),
        call_model=lambda *_: "布局图显示 AX-210 在左侧。",
    )
    assert any("AX-210" in c.text for c in parsed.chunks)
    assert any(c.locator["kind"] == "image_description" for c in parsed.chunks)


def test_pdf_content_image_gets_model_description(tmp_path):
    import pymupdf

    picture = tmp_path / "diagram.png"
    Image.new("RGB", (640, 400), "white").save(picture)
    path = tmp_path / "diagram.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((40, 50), "AX-210 drawing")
    page.insert_image(pymupdf.Rect(40, 80, 550, 410), filename=str(picture))
    doc.save(path)
    doc.close()
    parsed = enrich_pdf(
        parse_file(path, path.name), path, _settings(tmp_path),
        call_model=lambda *_: "图中标有 AX-210 及其连接位置。",
    )
    assert any("连接位置" in c.text and c.locator["page"] == 1 for c in parsed.chunks)
    assert not any("图片内容未解释" in warning for warning in parsed.warnings)


def test_zero_daily_limit_is_unlimited_even_after_historical_usage(tmp_path):
    from agico_kb.vision import _reserve, _store, _vision_call
    cfg = _settings(tmp_path, daily=0)
    conn = _store(cfg.storage_root)
    for _ in range(125):
        assert _reserve(conn, 0)
    calls = []
    assert _vision_call(conn, cfg, lambda: calls.append(1) or 'text') == 'text'
    assert calls == [1]
    assert conn.execute('SELECT count FROM daily_calls').fetchone()[0] == 126
    conn.close()


def test_vision_default_has_no_daily_quota(tmp_path, monkeypatch):
    monkeypatch.delenv('AGICO_KB_VISION_MAX_CALLS_PER_DAY', raising=False)
    assert Settings('unused', tmp_path).vision_max_calls_per_day == 0


def test_streaming_transcription_keeps_complete_token_output(monkeypatch):
    import json

    from agico_kb.vision import _call_model
    factory = httpx.Client
    # Simulate a healthy stream taking longer than the old four-minute cutoff.
    clock = iter(range(0, 10000, 300))
    monkeypatch.setattr('agico_kb.vision.time.monotonic', lambda: next(clock))
    def response(request):
        payload = json.loads(request.content)
        assert payload['stream'] is True
        events = [{'choices': [{'delta': {'content': '620 '}, 'finish_reason': None}]},
                  {'choices': [{'delta': {'content': 'mm'}, 'finish_reason': 'stop'}]}]
        body = ''.join('data: ' + json.dumps(e) + '\n\n' for e in events) + 'data: [DONE]\n\n'
        return httpx.Response(200, headers={'content-type': 'text/event-stream'}, text=body)
    monkeypatch.setattr('agico_kb.vision.httpx.Client', lambda **kwargs: factory(transport=httpx.MockTransport(response), **kwargs))
    assert _call_model('test', 'test', b'png') == '620 mm'


def test_incomplete_stream_is_not_accepted_as_finished_transcription(monkeypatch):
    from agico_kb.vision import _call_model
    factory = httpx.Client
    monkeypatch.setattr('agico_kb.vision.httpx.Client', lambda **kwargs: factory(transport=httpx.MockTransport(lambda _: httpx.Response(200, headers={'content-type': 'text/event-stream'}, text='data: {"choices":[{"delta":{"content":"partial"},"finish_reason":"length"}]}\n\ndata: [DONE]\n\n')), **kwargs))
    with pytest.raises(ValueError, match='输出未完成'):
        _call_model('test', 'test', b'png')


def test_provider_balance_error_is_explicit_and_terminal(monkeypatch):
    import json

    from agico_kb.vision import VisionQuotaError, _call_model
    factory = httpx.Client
    def response(request):
        assert json.loads(request.content)['max_tokens'] == 4096
        return httpx.Response(400, json={'error': {'code': 'insufficient_user_quota',
                                                  'message': 'credit insufficient balance'}})
    monkeypatch.setattr('agico_kb.vision.httpx.Client', lambda **kwargs: factory(
        transport=httpx.MockTransport(response), **kwargs))
    with pytest.raises(VisionQuotaError, match='余额不足') as error:
        _call_model('test', 'test', b'png')
    assert error.value.terminal


def test_official_glm_route_sends_image_and_selected_model(monkeypatch):
    import json

    from agico_kb.vision import _call_model
    factory = httpx.Client
    official = 'https://open.bigmodel.cn/api/paas/v4/chat/completions'
    monkeypatch.setattr('agico_kb.vision.API_URL', official)
    def response(request):
        assert str(request.url) == official
        payload = json.loads(request.content)
        assert payload['model'] == 'glm-5.3-flash'
        assert payload['reasoning_effort'] == 'low'
        assert payload['messages'][1]['content'][1]['type'] == 'image_url'
        return httpx.Response(200, json={'choices': [{'message': {'content': '620 mm'},
                                                     'finish_reason': 'stop'}]})
    monkeypatch.setattr('agico_kb.vision.httpx.Client', lambda **kwargs: factory(
        transport=httpx.MockTransport(response), **kwargs))
    assert _call_model('test', 'glm-5.3-flash', b'png') == '620 mm'


def test_connection_failover_uses_backup_key_then_returns_to_primary(tmp_path, monkeypatch):
    import json
    from dataclasses import replace

    from agico_kb.vision import PROMPT, _call_model, _describe
    primary = tmp_path/'primary.txt'
    backup = tmp_path/'backup.txt'
    primary.write_text('api=primary-secret\nmodel=glm-5.3-flash')
    backup.write_text('api=backup-secret\nmodel=glm-5.3-flashx')
    cfg = replace(_settings(tmp_path), vision_config_file=primary,
                  vision_fallback_config_file=backup, vision_fallback_model='glm-5.3-flashx')
    official = 'https://open.bigmodel.cn/api/paas/v4/chat/completions'
    monkeypatch.setattr('agico_kb.vision.API_URL', official)
    monkeypatch.setattr('agico_kb.vision.time.sleep', lambda _: None)
    factory = httpx.Client
    calls = []
    healthy = False
    def response(request):
        payload = json.loads(request.content)
        calls.append((str(request.url), payload['model'], request.headers['authorization']))
        if str(request.url) == official and not healthy:
            return httpx.Response(503, json={'error': {'message': 'unavailable'}})
        return httpx.Response(200, json={'choices': [{'message': {'content': '620 mm'}, 'finish_reason': 'stop'}]})
    monkeypatch.setattr('agico_kb.vision.httpx.Client', lambda **kwargs: factory(
        transport=httpx.MockTransport(response), **kwargs))
    buffer = BytesIO()
    Image.new('RGB', (40,40), 'white').save(buffer, format='PNG')
    for version in ['first', 'second']:
        assert _describe(buffer.getvalue(), cfg, version=version, prompt=PROMPT,
                        call_model=_call_model, screen=False) == '620 mm'
        healthy = True
    assert calls == [(official,'glm-5.3-flash','Bearer primary-secret')]*2 + [
        (cfg.vision_fallback_api_url,'glm-5.3-flashx','Bearer backup-secret'),
        (official,'glm-5.3-flash','Bearer primary-secret')]
