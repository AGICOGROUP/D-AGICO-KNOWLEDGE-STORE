def test_normal_processing_messages_do_not_alarm_users():
    from agico_kb.parse_notices import public_warnings

    messages = [
        'PDF 表格按原页保留；跨页表头和脚注需结合相邻页读取。',
        'OCR 识别文字未经人工核对；表格结构与图形语义可能不完整，关键数字请核对原文件。',
        '文本中的不可见空字符已替换为空格；原文件保留。',
        'PDF 5 页扫描正文采用 OCR 提取；未调用视觉 API。',
        '工作表“参数”保留公式与已有缓存结果，未重新计算公式。',
        'PDF 已生成 6 个视觉分区转写，0 页全部分区完成且未发现数字线索遗漏；细节请核对原文件。',
        'Word 文本框内容已收录（2 个文本框），版面位置与原文件不同。',
        'PDF 1 页视觉内容已生成 AI 转述；细节请核对原文件。',
        'PDF 第 1 页为扫描页，已 OCR 识别；关键数字请核对原文件。',
        '视觉分区处理已保存进度，排队继续处理。',
    ]
    assert public_warnings(messages) == []
    assert len(messages) == 10  # Presentation never rewrites stored diagnostics.


def test_actual_gaps_and_unknown_messages_are_never_hidden():
    from agico_kb.parse_notices import public_warnings

    messages = [
        'PDF 第 7、8 页数字或符号存在遗漏／冲突，尚未核定；请核对原图。',
        'Word 脚注／尾注未解析，内容可能缺少适用条件，请查看原文件。',
        '工作表“参数”公式结果未解析（原文件无缓存），请核对原文件。',
        'PDF 还有 2 页视觉内容未解释；请查看原文件。',
        '未知解析故障：内容可能缺失。',
        '视觉 API 暂时无响应，已保存进度并自动重试。',
    ]
    assert public_warnings(messages + messages) == messages


def test_drawing_scope_is_one_neutral_original_context_notice():
    from agico_kb.parse_notices import DRAWING_NOTICE, public_warnings

    messages = [
        f'PDF 第 {p} 页按图纸建立检索索引；图名、图号、说明和明细请结合原图核对，不转述全图尺寸。'
        for p in range(1, 24)
    ] + ['CAD 索引范围仅限原生文字和块属性；图形及精确尺寸关系请以原图为准。']
    assert public_warnings(messages) == [DRAWING_NOTICE]


def test_native_multipage_body_does_not_invent_table_limitation(tmp_path):
    import pymupdf

    from agico_kb.ingestion import parse_file

    path = tmp_path / 'fresh-body.pdf'
    with pymupdf.open() as doc:
        for i in range(3):
            doc.new_page().insert_text((60, 90), f'Equipment installation procedure step {i+1}.')
        doc.save(path)
    result = parse_file(path, path.name)
    assert result.status == 'ready'
    assert not any('表格' in w for w in result.warnings)


def test_verified_ocr_only_pdf_needs_no_vision_configuration(tmp_path):
    import pymupdf

    from agico_kb.config import Settings
    from agico_kb.ingestion import Parsed
    from agico_kb.vision import enrich_pdf

    path = tmp_path / 'fresh-scan.pdf'
    with pymupdf.open() as doc:
        page = doc.new_page()
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20))
        pix.clear_with(255)
        page.insert_image(page.rect, stream=pix.tobytes('png'))
        doc.save(path)
    parsed = Parsed(status='partial', warnings=[
        'OCR 识别文字未经人工核对；表格结构与图形语义可能不完整，关键数字请核对原文件。',
        'PDF 第 1 页为扫描页，已 OCR 识别；关键数字请核对原文件。'])
    parsed.add('Each component shall be inspected before installation.', kind='ocr', page=1,
               pdf_assessment={'route': 'ocr_text'})
    result = enrich_pdf(parsed, path, Settings('unused', tmp_path))
    assert result.status == 'ready'
    assert not any('未经' in w for w in result.warnings)

    parsed.status = 'partial'
    parsed.warnings.append('第 2-3 页解析未完成，该范围请核对原文件。')
    result = enrich_pdf(parsed, path, Settings('unused', tmp_path))
    assert result.status == 'partial'  # One successful OCR page cannot clear another page's failure.


def test_docx_notes_preserve_content_reference_and_table(tmp_path):
    import hashlib
    import zipfile

    from docx import Document
    from docx.oxml import parse_xml

    from agico_kb.ingestion import parse_file

    path = tmp_path / 'fresh-notes.docx'
    doc = Document()
    p = doc.add_paragraph('Maximum pressure 350 kPa')
    p._p.append(parse_xml('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnoteReference w:id="1"/></w:r>'))
    doc.save(path)
    with zipfile.ZipFile(path, 'a') as archive:
        for part, tag, value in [('footnotes', 'footnote', 'Only valid below 80 C.'),
                                 ('endnotes', 'endnote', 'Inspect after 24 hours.')]:
            archive.writestr(f'word/{part}.xml', f'''<w:{part} xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                <w:{tag} w:id="-1" w:type="separator"><w:p><w:r><w:t>IGNORE SEPARATOR</w:t></w:r></w:p></w:{tag}>
                <w:{tag} w:id="1"><w:p><w:r><w:t>{value}</w:t></w:r></w:p>
                  <w:tbl><w:tr><w:tc><w:p><w:r><w:t>Torque</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>330 Nm</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
                </w:{tag}></w:{part}>''')
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    parsed = parse_file(path, path.name)
    assert parsed.status == 'ready'
    assert any('[脚注 1]' in c.text for c in parsed.chunks)
    assert any(c.locator.get('footnote') == '1' and 'Only valid below 80 C.' in c.text for c in parsed.chunks)
    assert any(c.locator.get('endnote') == '1' and 'Inspect after 24 hours.' in c.text for c in parsed.chunks)
    assert any(c.locator.get('kind') == 'table' and 'Torque | 330 Nm' in c.text for c in parsed.chunks)
    assert not any('IGNORE SEPARATOR' in c.text for c in parsed.chunks)
    assert not any('未解析' in w for w in parsed.warnings)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_excel_text_shapes_are_searchable_with_sheet_and_anchor(tmp_path):
    import zipfile

    from openpyxl import Workbook

    from agico_kb.ingestion import parse_file

    path = tmp_path / 'fresh-shape.xlsx'
    book = Workbook()
    book.active.title = 'Parts'
    book.active.append(['Part', 'Quantity'])
    book.active.append(['Bolt', 12])
    book.save(path)
    with zipfile.ZipFile(path, 'a') as archive:
        archive.writestr('xl/worksheets/_rels/sheet1.xml.rels', '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>''')
        archive.writestr('xl/drawings/drawing1.xml', '''<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><xdr:twoCellAnchor><xdr:from><xdr:col>2</xdr:col><xdr:row>3</xdr:row></xdr:from><xdr:sp><xdr:txBody><a:bodyPr/><a:p><a:r><a:t>Use grade 8.8 bolts, torque 330 Nm.</a:t></a:r></a:p><a:p><a:r><a:t>Recheck after 24 hours.</a:t></a:r></a:p></xdr:txBody></xdr:sp></xdr:twoCellAnchor></xdr:wsDr>''')
    parsed = parse_file(path, path.name)
    assert parsed.status == 'ready'
    chunk = next(c for c in parsed.chunks if c.locator.get('kind') == 'drawing_text')
    assert chunk.locator['sheet'] == 'Parts'
    assert chunk.locator['cell'] == 'C4'
    assert 'torque 330 Nm.\nRecheck after 24 hours.' in chunk.text
    assert not any('绘图文字未解析' in w for w in parsed.warnings)


def test_numeric_check_ignores_identified_footer_but_keeps_same_number_in_body(tmp_path):
    import pymupdf

    from agico_kb.pdf_integrity import numeric_gaps
    from agico_kb.vision import _pdf_numeric_evidence

    path = tmp_path / 'fresh-numbered-procedure.pdf'
    with pymupdf.open() as doc:
        for i in range(3):
            page = doc.new_page()
            page.insert_text((70, 160), 'Allowed pressure 330 kPa, cycle count 1.')
            page.insert_text((70, 805), 'Contact 123456')
            page.insert_text((500, 805), str(i+1))
        doc.save(path)
    with pymupdf.open(path) as doc:
        evidence = _pdf_numeric_evidence(b'', doc[0], doc[0].rect)
    assert numeric_gaps(evidence, 'Allowed pressure 330 kPa, cycle count 1.') == []
    assert numeric_gaps(evidence, 'Allowed pressure 350 kPa, cycle count 1.') == ['330']


def test_raster_numeric_check_uses_footer_positions_not_just_number_values(tmp_path, monkeypatch):
    import io
    from types import SimpleNamespace

    import pymupdf
    from PIL import Image

    from agico_kb.vision import _pdf_numeric_evidence

    path = tmp_path / 'fresh-raster-footer.pdf'
    image = io.BytesIO()
    Image.new('RGB', (600, 900), 'white').save(image, format='PNG')
    with pymupdf.open() as doc:
        for _ in range(3):
            page = doc.new_page(width=600, height=900)
            page.insert_image(page.rect, stream=image.getvalue())
            page.insert_text((60, 870), 'Contact 123456')
        doc.save(path)
    boxes = [[[60, 180], [500, 180], [500, 205], [60, 205]],
             [[60, 850], [500, 850], [500, 880], [60, 880]]]
    output = SimpleNamespace(txts=['Serial 123456, pressure 330 kPa', 'Contact 123456'],
                             scores=[.99, .99], boxes=boxes)
    monkeypatch.setattr('agico_kb.ingestion._ocr_engine', lambda: lambda _: output)
    with pymupdf.open(path) as doc:
        evidence = _pdf_numeric_evidence(image.getvalue(), doc[0], doc[0].rect)
    assert evidence == ['Serial 123456, pressure 330 kPa']
