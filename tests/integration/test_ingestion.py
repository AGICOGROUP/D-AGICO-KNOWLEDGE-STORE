from uuid import uuid4

import pytest

from agico_kb.ingestion import parse_file


def test_docx_paragraph_table_and_real_source(tmp_path):
    from docx import Document

    path = tmp_path / (uuid4().hex + ".docx")
    doc = Document()
    doc.add_heading("安装条件", level=1)
    doc.add_paragraph("AX-210 允许温度 80 摄氏度，仅适用标准配置。")
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip(table.rows[0].cells, ["型号", "额定电压 (V)"], strict=True):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["AX-210", "220"], strict=True):
        cell.text = text
    doc.save(path)
    result = parse_file(path, path.name)
    assert result.status == "ready"
    assert any("80 摄氏度" in c.text and c.locator["section"] == "安装条件" for c in result.chunks)
    table_chunk = next(c for c in result.chunks if c.locator.get("table"))
    assert all(t in table_chunk.text for t in ["AX-210", "220", "额定电压 (V)"])
    assert "page" not in table_chunk.locator


def test_docx_stacked_title_characters_are_joined_into_one_chunk(tmp_path):
    """A stacked title arrives as one character per paragraph; it must not become four chunks."""
    from docx import Document

    path = tmp_path / (uuid4().hex + ".docx")
    doc = Document()
    doc.add_paragraph("300t/d回转窑活性石灰生产线")
    for character in "方案简介":
        doc.add_paragraph(character)
    doc.add_paragraph("本方案适用于菏泽双龙冶金机械有限公司的活性石灰生产线。")
    doc.save(path)
    result = parse_file(path, path.name)
    texts = [c.text for c in result.chunks]
    assert "方案简介" in texts
    assert not any(t in {"方", "案", "简", "介"} for t in texts)
    assert any("菏泽双龙" in t for t in texts)


def test_xlsx_preserves_formula_cache_distinction_and_coordinates(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / (uuid4().hex + ".xlsx")
    book = Workbook()
    sheet = book.active
    sheet.title = "参数"
    sheet.append(["型号", "数量", "单价", "合计"])
    sheet.append(["AX-210", 2, 150, "=B2*C2"])
    book.save(path)
    result = parse_file(path, path.name)
    text = "\n".join(c.text for c in result.chunks)
    assert all(v in text for v in ["AX-210", "150", "=B2*C2", "缓存结果缺失"])
    assert result.chunks[0].locator["sheet"] == "参数"
    assert result.chunks[0].locator["range"] == "A1:D2"
    assert any("公式" in w for w in result.warnings)


def test_pptx_slide_source(tmp_path):
    from pptx import Presentation

    path = tmp_path / (uuid4().hex + ".pptx")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "合成业务背景"
    slide.placeholders[1].text = "客户文件应注明产品型号和单位。"
    slides.save(path)
    result = parse_file(path, path.name)
    assert any("产品型号和单位" in c.text and c.locator["slide"] == 1 for c in result.chunks)


@pytest.mark.parametrize("case", ["populated", "empty", "oversized"])
@pytest.mark.parametrize("stored_suffix", [".xlsx", ".blob"])
def test_xlsx_without_optional_dimension_metadata(tmp_path, case, stored_suffix):
    import re
    import zipfile

    from openpyxl import Workbook

    original = tmp_path / (uuid4().hex + ".xlsx")
    path = tmp_path / (uuid4().hex + stored_suffix)
    book = Workbook()
    if case == "populated":
        book.active.append(["设备", "数量", "合计"])
        book.active.append(["风机", 2, "=B2*3"])
    elif case == "oversized":
        book.active.cell(1001, 1000, "too large")
    book.save(original)
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(path, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename.startswith("xl/worksheets/"):
                data = re.sub(rb"<dimension\b[^>]*/>", b"", data)
            target.writestr(entry, data)
    if case == "oversized":
        with pytest.raises(ValueError, match="保护上限"):
            parse_file(path, original.name)
        return
    result = parse_file(path, original.name)
    if case == "empty":
        assert not result.chunks
    else:
        assert result.chunks[0].locator["range"] == "A1:C2"
        assert all(s in result.chunks[0].text for s in ["风机", "=B2*3", "缓存结果缺失"])


def test_native_pdf_page_and_units(tmp_path):
    import pymupdf

    path = tmp_path / (uuid4().hex + ".pdf")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((60, 80), "AX-210 max temperature: 80 C. Standard configuration only.")
    doc.save(path)
    doc.close()
    result = parse_file(path, path.name)
    assert any("80 C" in c.text and c.locator["page"] == 1 for c in result.chunks)
    assert not any(c.locator.get("ocr") for c in result.chunks)


def test_scan_pdf_ocr_preserves_source_and_reports_uncertainty(tmp_path):
    import pymupdf

    path = tmp_path / (uuid4().hex + ".pdf")
    source = pymupdf.open()
    p = source.new_page(width=600, height=240)
    p.insert_text((30, 100), "AX-210 rated voltage 220 V", fontsize=24)
    pix = p.get_pixmap(dpi=150)
    scanned = pymupdf.open()
    page = scanned.new_page(width=600, height=240)
    page.insert_image(page.rect, stream=pix.tobytes("png"))
    scanned.save(path)
    source.close()
    scanned.close()
    result = parse_file(path, path.name)
    assert any(
        "220" in c.text and c.locator.get("ocr") and c.locator["page"] == 1 for c in result.chunks
    )
    assert any("OCR" in w for w in result.warnings)


@pytest.mark.parametrize("suffix", [".step", ".dwg", ".zip"])
def test_unsupported_attachment_is_stored_not_silently_parsed(tmp_path, suffix):
    # .doc/.ppt/.xls now go through the pure-Python fallback instead of stored_only,
    # so they are no longer part of the opaque-attachment set.
    path = tmp_path / (uuid4().hex + suffix)
    path.write_bytes(b"new synthetic opaque attachment")
    result = parse_file(path, path.name)
    assert result.status == "stored_only" and result.chunks == [] and result.warnings


def test_markdown_keeps_long_semantic_unit_and_footnotes(tmp_path):
    path = tmp_path / (uuid4().hex + ".md")
    text = "# 条件\n\n" + ("适用于标准配置。" * 700) + "\n\n注：不适用于防爆环境。"
    path.write_text(text, encoding="utf-8")
    result = parse_file(path, path.name)
    assert "不适用于防爆环境" in "\n".join(c.text for c in result.chunks)
    assert any(len(c.text) > 4000 for c in result.chunks)


def test_pdf_text_block_crossing_table_boundary_keeps_outside_condition(tmp_path):
    import pymupdf

    path = tmp_path / (uuid4().hex + ".pdf")
    doc = pymupdf.open()
    page = doc.new_page()
    for y in (80, 110, 140):
        page.draw_line((50, y), (300, y))
    for x in (50, 175, 300):
        page.draw_line((x, 80), (x, 140))
    page.insert_text((60, 100), "Model")
    page.insert_text((190, 100), "Voltage")
    page.insert_text((60, 125), "AX-210")
    page.insert_text((190, 125), "220 V\nSTANDARD ONLY\nNOT EXPLOSION PROOF", fontsize=11)
    doc.save(path)
    doc.close()
    result = parse_file(path, path.name)
    text = "\n".join(c.text for c in result.chunks)
    assert "220 V" in text and "NOT EXPLOSION PROOF" in text


def test_docx_textbox_content_is_captured_and_not_silently_dropped(tmp_path):
    from docx import Document
    from docx.oxml import parse_xml

    path = tmp_path / (uuid4().hex + ".docx")
    doc = Document()
    paragraph = doc.add_paragraph("AX-210 max temperature 80 C.")
    paragraph._p.append(
        parse_xml(
            '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>ONLY STANDARD CONFIGURATION</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>'
        )
    )
    doc.save(path)
    result = parse_file(path, path.name)
    text = "\n".join(c.text for c in result.chunks)
    assert "ONLY STANDARD CONFIGURATION" in text
    assert any(c.locator.get("textbox") for c in result.chunks)
    # Everything in the file is readable now, so it is ready rather than partial, but the reader
    # is still told the box's position differs from the original layout.
    assert result.status == "ready"
    assert any("文本框" in w for w in result.warnings)


def test_docx_textbox_table_is_kept_as_a_table_once(tmp_path):
    """A boxed table must be indexed with its rows, and mc:AlternateContent must not double it."""
    from docx import Document
    from docx.oxml import parse_xml

    path = tmp_path / (uuid4().hex + ".docx")
    doc = Document()
    doc.add_paragraph("除臭措施见下表。")
    box = (
        "<w:txbxContent>"
        "<w:tbl>"
        "<w:tr><w:tc><w:p><w:r><w:t>场所</w:t></w:r></w:p></w:tc>"
        "<w:tc><w:p><w:r><w:t>除臭措施</w:t></w:r></w:p></w:tc></w:tr>"
        "<w:tr><w:tc><w:p><w:r><w:t>垃圾卸料大厅</w:t></w:r></w:p></w:tc>"
        "<w:tc><w:p><w:r><w:t>植物除臭剂喷洒</w:t></w:r></w:p></w:tc></w:tr>"
        "</w:tbl></w:txbxContent>"
    )
    paragraph = doc.add_paragraph("")
    paragraph._p.append(
        parse_xml(
            '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><v:shape><v:textbox>'
            + box
            + "</v:textbox></v:shape></w:pict>"
        )
    )
    doc.save(path)
    result = parse_file(path, path.name)
    boxed = [c for c in result.chunks if c.locator.get("textbox")]
    assert len(boxed) == 1
    assert "垃圾卸料大厅 | 植物除臭剂喷洒" in boxed[0].text
    assert boxed[0].locator["kind"] == "table"


def test_docx_content_control_paragraphs_are_indexed(tmp_path):
    """A field-generated table of contents lives in w:sdt; skipping it loses the whole TOC."""
    from docx import Document
    from docx.oxml import parse_xml

    path = tmp_path / (uuid4().hex + ".docx")
    doc = Document()
    doc.add_paragraph("第一章 工程设计")
    body = doc.element.body
    body.append(
        parse_xml(
            '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:sdtContent>"
            "<w:p><w:r><w:t>1.4 设计规模Escala de diseño</w:t></w:r><w:r><w:tab/></w:r>"
            "<w:r><w:t>- 2 -</w:t></w:r></w:p>"
            "</w:sdtContent></w:sdt>"
        )
    )
    doc.save(path)
    result = parse_file(path, path.name)
    text = "\n".join(c.text for c in result.chunks)
    assert "1.4 设计规模" in text
    assert any(c.locator.get("control") for c in result.chunks)
    assert any("内容控件内容已收录" in w for w in result.warnings)


def test_docx_header_table_missing_from_body_is_disclosed(tmp_path):
    from docx import Document
    from docx.shared import Inches

    path = tmp_path / (uuid4().hex + ".docx")
    doc = Document()
    doc.add_paragraph("AX-210 80 C.")
    table = doc.sections[0].header.add_table(rows=1, cols=1, width=Inches(4))
    table.cell(0, 0).text = "仅标准配置有效"
    doc.save(path)
    result = parse_file(path, path.name)
    assert result.status == "partial" and any("页眉" in w for w in result.warnings)
